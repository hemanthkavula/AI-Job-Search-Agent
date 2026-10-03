from pathlib import Path
import json

ROOT=Path(__file__).resolve().parents[1]

def replace_once(path,old,new):
    p=ROOT/path; text=p.read_text(encoding='utf-8')
    n=text.count(old)
    if n!=1: raise RuntimeError(f'{path}: expected one match, found {n}: {old[:120]!r}')
    p.write_text(text.replace(old,new,1),encoding='utf-8')

# -----------------------------------------------------------------------------
# 1. Uploaded 2026-10-03 PDF becomes the master factual/content baseline.
# -----------------------------------------------------------------------------
profile_path=ROOT/'data/candidate_profile.json'
profile=json.loads(profile_path.read_text(encoding='utf-8'))
profile['master_resume_reference']={
    'filename':'Hemanth_Kavula_Senior_Data_Engineer_Resume.pdf',
    'effective_date':'2026-10-03',
    'layout_version':'master-2026-10-03',
    'pages':2,
    'bullet_counts':{'Fidelity Investments':10,'Cigna Healthcare':8,'Target Corporation':8},
    'style':{
        'font':'Calibri','name_pt':18,'headline_pt':13,'contact_pt':11,
        'section_heading_pt':12,'summary_pt':11,'body_pt':10,
        'accent_hex':'1F4E79','selective_phrase_bold':True,
        'roles_and_responsibilities_label':True,'environment_line':True,
        'right_aligned_dates':True,'two_summary_paragraphs':True
    }
}
profile['summary_source']=[
    'Senior Data Engineer with 5+ years of experience building high-scale data platforms and real-time pipelines processing millions of transactions and large-volume datasets across financial services, healthcare, and retail domains. Expert in Python, SQL, and PySpark with deep experience in Apache Spark, Kafka, and Databricks for distributed data processing and streaming systems.',
    'Proven track record of architecting cloud-native data lakes and warehouses on AWS and Azure (S3, Glue, EMR, Redshift, Snowflake, ADF, ADLS Gen2), delivering reliable, analytics-ready datasets for business-critical use cases including risk analytics, clinical insights, and retail operations. Strong expertise in dimensional modeling (Star Schema, SCD Type 2), medallion architecture, and performance optimization, consistently improving pipeline efficiency, scalability, and data quality.'
]
profile['summary_emphasis']=[
    'Senior Data Engineer with 5+ years of experience',
    'financial services, healthcare, and retail domains',
    'Python, SQL, and PySpark',
    'Apache Spark, Kafka, and Databricks',
    'AWS and Azure (S3, Glue, EMR, Redshift, Snowflake, ADF, ADLS Gen2)',
    'dimensional modeling (Star Schema, SCD Type 2)',
    'medallion architecture',
    'improving pipeline efficiency, scalability, and data quality'
]
profile['skill_categories']={
    'Programming Languages':['Python','SQL','Scala'],
    'Cloud Platforms (AWS)':['AWS Glue','Amazon EMR','Amazon S3','Amazon Redshift','Lambda','Kinesis'],
    'Cloud Platforms (Azure)':['Data Factory','Synapse Analytics','Azure Data Lake Storage Gen2','Event Hub'],
    'Big Data & Data Processing':['Apache Spark','PySpark','Databricks','Delta Lake','Apache Kafka'],
    'Data Warehousing':['Snowflake','Amazon Redshift','Azure Synapse Analytics'],
    'Databases':['Oracle','SQL Server','PostgreSQL','MySQL','MongoDB','DynamoDB'],
    'Data Integration & Orchestration':['Apache Airflow','AWS Glue','Azure Data Factory','dbt'],
    'Data Modeling':['Star Schema','Snowflake Schema','Dimensional Modeling','Slowly Changing Dimensions'],
    'DevOps & CI/CD':['Terraform','Docker','Jenkins','GitLab CI/CD','Git'],
    'Data Quality & Governance':['Great Expectations','AWS Glue Data Catalog','Azure Purview'],
    'Data Visualization & BI':['Tableau','Power BI','Looker'],
    'Methodologies':['Agile/Scrum','DevOps','DataOps','CI/CD Best Practices']
}
# Keep the hidden verified inventory synchronized with the visible new master.
profile['skills']=list(dict.fromkeys(v for values in profile['skill_categories'].values() for v in values))

exp={x['company']:x for x in profile['experience']}
exp['Fidelity Investments']['evidence']=[
    'Designed and owned AWS-based data pipelines using Glue and EMR (PySpark), processing ~2.5–3M daily trading events and ~400–500GB of market data, enabling risk monitoring and regulatory reporting for investment and compliance teams.',
    'Built real-time streaming pipelines using Kafka and Kinesis, ingesting trade lifecycle events with sub-minute latency to improve visibility for compliance monitoring systems.',
    'Modeled and delivered Snowflake data warehouse solutions (Star Schema, SCD Type 2) supporting portfolio analytics and regulatory reporting used by risk and investment stakeholders.',
    'Architected and maintained an S3-based data lake (Delta, medallion architecture) supporting curated datasets consumed by analysts and quantitative teams.',
    'Developed modular dbt transformation models and automated data quality tests in Snowflake, creating reusable analytics datasets for risk and investment reporting.',
    'Established data quality validation using Great Expectations (70+ rules), improving data reliability and reducing downstream reporting inconsistencies.',
    'Enabled AI/ML-driven analytics by building scalable Python and PySpark data pipelines that curated trading and market datasets for anomaly detection, risk analysis, and predictive modeling use cases.',
    'Optimized Spark workloads through partitioning and join strategies, reducing pipeline runtime by ~30% and lowering EMR compute costs.',
    'Automated infrastructure provisioning and deployments using Terraform and Jenkins, ensuring consistent and reliable releases across environments.',
    'Enabled analytics-ready datasets for portfolio managers and quant teams, supporting risk exposure analysis and trading performance insights.'
]
exp['Fidelity Investments']['bullet_emphasis']=[
    ['~2.5–3M daily trading events and ~400–500GB of market data'],['Kafka and Kinesis'],
    ['Snowflake data warehouse solutions (Star Schema, SCD Type 2)'],['S3-based data lake (Delta, medallion architecture)'],
    ['dbt transformation models'],['Great Expectations (70+ rules)'],
    ['AI/ML-driven analytics','Python and PySpark data pipelines'],['~30%','EMR compute costs'],
    ['Terraform and Jenkins'],['analytics-ready datasets for portfolio managers and quant teams']
]
exp['Cigna Healthcare']['bullet_emphasis']=[
    ['claims, eligibility, and provider data for ~2–3M members'],['Databricks'],
    ['ADLS Gen2 data lake (layered/medallion architecture)'],['Snowflake data warehouse solutions'],
    ['Event Hub and Structured Streaming'],['~25–30%'],['Azure Purview'],['operations and clinical teams']
]
exp['Target Corporation']['bullet_emphasis']=[
    ['~120–180GB/day of retail sales and inventory data'],['product, pricing, and order data'],
    ['Redshift data warehouse models (Star Schema)'],['near real-time POS data'],['Python and SQL'],
    ['PostgreSQL/MySQL'],['merchandising and operations teams'],['data-driven decision-making for retail operations']
]
profile_path.write_text(json.dumps(profile,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')

# -----------------------------------------------------------------------------
# 2. Reference formatter: reproduce the uploaded PDF's visual grammar.
# -----------------------------------------------------------------------------
formatter='''from __future__ import annotations
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.resume_generator import ROOT, safe_name, clean_company_name
from datetime import datetime
from zoneinfo import ZoneInfo
import re

ACCENT=RGBColor(31,78,121); GRAY=RGBColor(89,89,89); LINK=RGBColor(5,99,193); BLUE="1F4E79"
MASTER_BULLET_COUNTS={"Fidelity Investments":10,"Cigna Healthcare":8,"Target Corporation":8}

def _run(r,size=10,bold=False,color=None,underline=False):
    r.font.name="Calibri";r.font.size=Pt(size);r.bold=bold;r.underline=underline
    if color:r.font.color.rgb=color
    return r

def _compact(p,before=0,after=0,line=1.0):
    p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=line

def _section(doc,text,after=5):
    p=doc.add_paragraph();_compact(p,10,after);p.paragraph_format.keep_with_next=True;_run(p.add_run(text),12,True,ACCENT)
    pPr=p._p.get_or_add_pPr();pBdr=OxmlElement("w:pBdr");bottom=OxmlElement("w:bottom")
    bottom.set(qn("w:val"),"single");bottom.set(qn("w:sz"),"8");bottom.set(qn("w:space"),"2");bottom.set(qn("w:color"),BLUE);pBdr.append(bottom);pPr.append(pBdr)
    return p

def _emphasis_runs(p,text,phrases,size=10,color=None):
    text=str(text or "")
    candidates=[str(x).strip() for x in (phrases or []) if str(x).strip()]
    seen=set();selected=[]
    for phrase in sorted(candidates,key=len,reverse=True):
        key=phrase.casefold()
        if key not in seen and key in text.casefold():seen.add(key);selected.append(phrase)
    if not selected:
        _run(p.add_run(text),size,False,color);return
    pattern=re.compile("("+"|".join(re.escape(x) for x in selected)+")",re.I)
    lookup={x.casefold() for x in selected}
    for part in pattern.split(text):
        if part:_run(p.add_run(part),size,part.casefold() in lookup,color)

def _summary_parts(text):
    text=str(text or "").strip();parts=[x.strip() for x in re.split(r"\\n\\s*\\n",text) if x.strip()]
    if len(parts)>=2:return parts[:2]
    sentences=[x.strip() for x in re.split(r"(?<=[.!?])\\s+",text) if x.strip()]
    if len(sentences)>=4:return [" ".join(sentences[:2])," ".join(sentences[2:])]
    if len(sentences)==3:return [sentences[0]," ".join(sentences[1:])]
    return [text] if text else []

def _all_skill_terms(generated):
    out=[]
    for vals in (generated.get("skills") or {}).values():out.extend(str(x) for x in (vals or []))
    return list(dict.fromkeys(out))

def _metric_terms(text):
    patterns=[r"~?\\d+(?:\\.\\d+)?\\s*[–-]\\s*\\d+(?:\\.\\d+)?\\s*(?:GB/day|GB|TB|%|M|million)?",r"~?\\d+(?:\\.\\d+)?\\s*(?:GB/day|GB|TB|%|M|million)"]
    hits=[]
    for pat in patterns:hits.extend(m.group(0) for m in re.finditer(pat,text,re.I))
    return list(dict.fromkeys(hits))

def _fallback_emphasis(text,generated,max_terms=2):
    hits=[]
    for term in sorted(_all_skill_terms(generated),key=len,reverse=True):
        if len(term)>=3 and re.search(r"(?<![A-Za-z0-9])"+re.escape(term)+r"(?![A-Za-z0-9])",text,re.I):hits.append(term)
    hits=_metric_terms(text)+hits
    return list(dict.fromkeys(hits))[:max_terms]

def _company_header(doc,base):
    p=doc.add_paragraph();_compact(p,7,0);p.paragraph_format.keep_with_next=True;p.paragraph_format.tab_stops.add_tab_stop(Inches(7.45),WD_TAB_ALIGNMENT.RIGHT)
    _run(p.add_run(base["company"]),12,True);_run(p.add_run(" | "+base.get("location","")),10);_run(p.add_run("\\t"+base["dates"]),10,False,GRAY)
    title=doc.add_paragraph();_compact(title,0,1);title.paragraph_format.keep_with_next=True;_run(title.add_run(base["title"]),11,True,ACCENT)
    roles=doc.add_paragraph();_compact(roles,0,1);roles.paragraph_format.keep_with_next=True;_run(roles.add_run("Roles & Responsibilities:"),10,True)
    return p,title,roles

def canonical_resume_title(title):
    raw=re.sub(r"\\s+"," ",str(title or "")).strip()
    m=re.search(r"(?i)\\b(?:(principal|staff|lead|senior|sr\\.?|junior|jr\\.?)\\s+)?data\\s+engineer(?:ing)?\\b",raw)
    if m:
        level=(m.group(1) or "").lower().rstrip(".");level={"sr":"Senior","jr":"Junior"}.get(level,level.title())
        core="Data Engineering" if re.search(r"(?i)data\\s+engineering",m.group(0)) else "Data Engineer";return f"{level} {core}".strip()
    cleaned=re.sub(r"(?i)^\\s*(?:immediate interviews?|urgent(?: hiring)?|hiring now)\\s*[-:|]\\s*","",raw)
    cleaned=re.split(r"\\s+(?:[-|/]\\s*)(?=(?:airflow|dbt|kubernetes|openshift|aws|azure|gcp|snowflake|databricks|hybrid|remote|onsite|on-site)\\b)",cleaned,1,flags=re.I)[0]
    return cleaned.strip(" -|:/") or "Senior Data Engineer"

def render_llm_resume(job,profile,generated,output_dir="generated/resumes"):
    """Render every base/JD-specific resume in the uploaded 2026-10-03 master style."""
    doc=Document();s=doc.sections[0];s.top_margin=Inches(.55);s.bottom_margin=Inches(.42);s.left_margin=Inches(.5);s.right_margin=Inches(.5)
    doc.styles["Normal"].font.name="Calibri";doc.styles["Normal"].font.size=Pt(10);doc.styles["Normal"].paragraph_format.space_after=Pt(0)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;_compact(p,0,2.5);_run(p.add_run(profile["name"]),18,True,ACCENT)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;_compact(p,0,3);_run(p.add_run(canonical_resume_title(job.title or profile.get("headline","Senior Data Engineer"))),13,False,ACCENT)
    c=profile.get("contact",{});p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;_compact(p,0,7)
    vals=[x for x in [c.get("phone"),c.get("email"),c.get("linkedin")] if x]
    for i,v in enumerate(vals):
        if i:_run(p.add_run(" | "),11)
        link=("@" in v or "linkedin" in v.lower());_run(p.add_run(v),11,False,LINK if link else None,link)

    _section(doc,"PROFESSIONAL SUMMARY",4)
    summary_emphasis=generated.get("summary_emphasis") or profile.get("summary_emphasis") or []
    parts=_summary_parts(generated.get("summary",""))
    for i,part in enumerate(parts):
        p=doc.add_paragraph();_compact(p,0,7 if i==0 and len(parts)>1 else 3,1.05)
        phrases=[x for x in summary_emphasis if str(x).casefold() in part.casefold()] or _fallback_emphasis(part,generated,5)
        _emphasis_runs(p,part,phrases,11)

    _section(doc,"TECHNICAL SKILLS",7)
    for label,vals in (generated.get("skills") or {}).items():
        p=doc.add_paragraph();_compact(p,0,.2,1.0);_run(p.add_run(str(label)+": "),10,True);_run(p.add_run(", ".join(str(v) for v in vals)),10)

    _section(doc,"PROFESSIONAL EXPERIENCE",4)
    expected={x["company"]:x for x in profile["experience"]}
    for item in generated.get("experience",[]):
        base=expected.get(item.get("company"))
        if not base:continue
        company_p,title_p,roles_p=_company_header(doc,base)
        first_bullet=None;emphasis_rows=item.get("bullet_emphasis") or base.get("bullet_emphasis") or []
        limit=(profile.get("master_resume_reference",{}).get("bullet_counts",{}) or MASTER_BULLET_COUNTS).get(base["company"],8)
        for idx,line in enumerate((item.get("bullets") or [])[:limit]):
            p=doc.add_paragraph(style="List Bullet");p.paragraph_format.left_indent=Inches(.50);p.paragraph_format.first_line_indent=Inches(-.18);p.paragraph_format.keep_together=True;_compact(p,0,3,1.0)
            # The master intentionally continues Fidelity bullets 9-10 on page two.
            if base["company"]=="Fidelity Investments" and idx==8:p.paragraph_format.page_break_before=True
            phrases=emphasis_rows[idx] if idx<len(emphasis_rows) and isinstance(emphasis_rows[idx],list) else []
            _emphasis_runs(p,str(line).strip(),phrases or _fallback_emphasis(str(line),generated,2),10)
            if first_bullet is None:first_bullet=p
        env=doc.add_paragraph();_compact(env,1,8);_run(env.add_run("Environment: "),10,True);_run(env.add_run(base.get("environment","")),10)
        if first_bullet is None:roles_p.paragraph_format.keep_with_next=False

    _section(doc,"EDUCATION",4)
    for e in profile["education"]:
        p=doc.add_paragraph();_compact(p,0,1);p.paragraph_format.keep_with_next=True;_run(p.add_run(e["degree"]),11,True)
        p=doc.add_paragraph();_compact(p,0,0);p.paragraph_format.tab_stops.add_tab_stop(Inches(7.45),WD_TAB_ALIGNMENT.RIGHT);_run(p.add_run(f"{e['school']} | {e['location']}"),10);_run(p.add_run("\\t"+f"{e['start']} – {e['end']}"),10,False,GRAY)

    root=ROOT/output_dir;root.mkdir(parents=True,exist_ok=True);pattern=profile.get("output",{}).get("resume_filename_pattern","Hemanth_Kavula_{Company}_{JobTitle}")
    clean_title=canonical_resume_title(job.title);clean_company=clean_company_name(job.company)
    stem=pattern.replace("{Company}",safe_name(clean_company)).replace("{JobTitle}",safe_name(clean_title))
    timestamp=datetime.now(ZoneInfo("America/New_York")).strftime("%Y%m%d_%H%M%S")
    job_dir=root/f"{safe_name(clean_company)}_{safe_name(clean_title)}_{timestamp}";job_dir.mkdir(parents=True,exist_ok=True)
    path=job_dir/f"{stem}.docx";doc.save(path);return str(path)
'''
(ROOT/'app/reference_resume_formatter.py').write_text(formatter,encoding='utf-8')

# -----------------------------------------------------------------------------
# 3. Base-resume rendering preserves the master paragraphs and emphasis metadata.
# -----------------------------------------------------------------------------
replace_once('app/batch_prepare.py','"summary": " ".join(profile.get("summary_source") or []),','"summary": "\\n\\n".join(profile.get("summary_source") or []),\n        "summary_emphasis": list(profile.get("summary_emphasis") or []),')
replace_once('app/batch_prepare.py','{"company": row.get("company"), "bullets": list(row.get("evidence") or [])}','{"company": row.get("company"), "bullets": list(row.get("evidence") or []), "bullet_emphasis": list(row.get("bullet_emphasis") or [])}')

# -----------------------------------------------------------------------------
# 4. Tailored-resume contract: same master density/visual emphasis, JD-specific words.
# -----------------------------------------------------------------------------
writer=ROOT/'app/llm_resume_writer.py';text=writer.read_text(encoding='utf-8')
repls={
'Use exactly 8 Fidelity bullets, 7 Cigna bullets, and 6 Target bullets.':'Use exactly 10 Fidelity bullets, 8 Cigna bullets, and 8 Target bullets, matching the current uploaded master resume.',
'The master resume is a FIXED-FACTS SOURCE ONLY: use it only for candidate identity/contact details, employer names, job titles, locations, employment dates, and education. Do NOT use the master resume\'s Technical Skills, Professional Summary, or existing experience bullet content as the content source or ceiling for a new resume. Generate the new Professional Summary, Technical Skills, and all Professional Experience bullets from the CURRENT COMPLETE JOB DESCRIPTION, while preserving the fixed employment chronology and employer/domain context.':'The uploaded master resume is the candidate truth baseline. Preserve its identity/contact details, employers, titles, locations, dates, education, employer domains, approved environments, evidence boundaries, and approved metrics. Generate JD-specific Summary, Technical Skills, and Experience wording from the CURRENT COMPLETE JOB DESCRIPTION, but keep every claim interview-defensible against the master evidence and approved capabilities. Rewrite and re-emphasize; do not mechanically copy the master bullets.',
'The master profile is only the source of fixed identity, chronology, employers, roles, locations, dates, and education. Ignore its prior skills and bullet wording when generating tailored content.':'The master profile is the factual evidence baseline as well as the source of fixed identity and chronology. Tailor toward the JD without inventing employer/domain facts, unsupported metrics, or unsupported specific accomplishments; master wording may be rewritten when the same supported fact is expressed more effectively for the JD.',
'STRICT METRIC RULE: Fidelity may have at most 2 metric-bearing bullets, Cigna at most 2, and Target must contain ZERO numeric scale, percentage, volume, latency, count, or performance metrics.':'STRICT METRIC RULE: Fidelity may have at most 2 metric-bearing bullets, Cigna at most 2, and Target may have at most 1 metric-bearing bullet. Only metrics explicitly approved by the current master evidence may be used; Target\'s approved metric is ~120–180GB/day of retail sales and inventory data.',
'Before returning JSON, silently self-check: exact 8/7/6 bullet counts;':'Before returning JSON, silently self-check: exact 10/8/8 bullet counts;',
'"fidelity_metric_bullets_max":2,"cigna_metric_bullets_max":2,"target_metric_bullets_max":0':'"fidelity_metric_bullets_max":2,"cigna_metric_bullets_max":2,"target_metric_bullets_max":1',
'"fidelity_bullets":8,"cigna_bullets":7,"target_bullets":6':'"fidelity_bullets":10,"cigna_bullets":8,"target_bullets":8',
'"summary":"3 concise recruiter-friendly sentences aligned to the target title and strongest supported JD requirements",':'"summary":"2 concise paragraphs separated by a blank line, normally 4 concise sentences total, matching the density of the uploaded master while aligning to the target title and strongest supported JD requirements",\n        "summary_emphasis":["4-8 concise phrases copied exactly from summary that should appear bold, prioritizing target title, domain phrase, core technologies/architecture and strongest outcome phrase"],',
'"experience":[{"company":"exact employer","title":"exact title","dates":"exact dates","bullets":["strong JD-specific bullet that preserves employer/domain facts and does not invent a false specific accomplishment or metric"]}],':'"experience":[{"company":"exact employer","title":"exact title","dates":"exact dates","bullets":["strong JD-specific bullet, normally <=32 words, that preserves employer/domain facts and does not invent a false specific accomplishment or metric"],"bullet_emphasis":[["1-2 concise exact phrases from the corresponding bullet to bold, usually the most important metric, technology+system phrase, architecture phrase, or business-data phrase"]]}],'
}
for old,new in repls.items():
    if text.count(old)!=1:raise RuntimeError(f'llm_resume_writer.py replacement count {text.count(old)} for {old[:90]!r}')
    text=text.replace(old,new,1)
needle='Return valid JSON only with keys summary, skills, experience, and education.'
insert='''VISUAL EMPHASIS CONTRACT: Match the uploaded master resume's selective bolding. Return summary_emphasis as exact substrings from summary and bullet_emphasis as a list aligned one-for-one with each employer's bullets. Use 1-2 bold phrases per bullet, not whole bullets and not random keyword stuffing. Favor meaningful technology+system phrases, approved metrics, architecture/data phrases, and stakeholder/domain phrases. Emphasis metadata is formatting-only and must never introduce text that is not already present in the summary/bullet.\nReturn valid JSON only with keys summary, summary_emphasis, skills, experience, and education.'''
if text.count(needle)!=1:raise RuntimeError('LLM return schema anchor not found')
text=text.replace(needle,insert,1)
writer.write_text(text,encoding='utf-8')

# -----------------------------------------------------------------------------
# 5. Audit/fallback counts and approved metrics follow the new master.
# -----------------------------------------------------------------------------
replace_once('app/ats_audit.py','EXPECTED_COUNTS={"Fidelity Investments":8,"Cigna Healthcare":7,"Target Corporation":6}','EXPECTED_COUNTS={"Fidelity Investments":10,"Cigna Healthcare":8,"Target Corporation":8}')
replace_once('app/ats_audit.py','METRIC_LIMITS={"Fidelity Investments":2,"Cigna Healthcare":2,"Target Corporation":0}','METRIC_LIMITS={"Fidelity Investments":2,"Cigna Healthcare":2,"Target Corporation":1}')
replace_once('app/ats_audit.py','"Cigna Healthcare":[r"2\\s*[–-]\\s*3\\s*(?:m|million)",r"25\\s*[–-]\\s*30\\s*%"],"Target Corporation":[]','"Cigna Healthcare":[r"2\\s*[–-]\\s*3\\s*(?:m|million)",r"25\\s*[–-]\\s*30\\s*%"],"Target Corporation":[r"120\\s*[–-]\\s*180\\s*gb\\s*/?\\s*day"]')

rg=ROOT/'app/resume_generator.py';rt=rg.read_text(encoding='utf-8')
rt=rt.replace('# Exact bullet counts: Fidelity 8, Cigna 7, Target 6.','# Exact bullet counts: Fidelity 10, Cigna 8, Target 8 (current master resume).')
old='{"Fidelity Investments":8,"Cigna Healthcare":7,"Target Corporation":6}'
if rt.count(old)<2:raise RuntimeError(f'resume_generator.py expected >=2 bullet-count dictionaries, found {rt.count(old)}')
rt=rt.replace(old,'{"Fidelity Investments":10,"Cigna Healthcare":8,"Target Corporation":8}')
rg.write_text(rt,encoding='utf-8')

# -----------------------------------------------------------------------------
# 6. Existing layout regression understands the Roles row; add master regression.
# -----------------------------------------------------------------------------
test=ROOT/'tests/test_reference_resume_formatter_layout.py';tt=test.read_text(encoding='utf-8')
old='''    company=paragraphs[idx]; title=paragraphs[idx+1]; first=paragraphs[idx+2]\n    assert company.paragraph_format.keep_with_next is True\n    assert title.paragraph_format.keep_with_next is True\n    assert first.paragraph_format.keep_together is True\n    assert first.paragraph_format.keep_with_next is not True\n'''
new='''    company=paragraphs[idx]; title=paragraphs[idx+1]; roles=paragraphs[idx+2]; first=paragraphs[idx+3]\n    assert company.paragraph_format.keep_with_next is True\n    assert title.paragraph_format.keep_with_next is True\n    assert roles.text=="Roles & Responsibilities:"\n    assert roles.paragraph_format.keep_with_next is True\n    assert first.paragraph_format.keep_together is True\n    assert first.paragraph_format.keep_with_next is not True\n'''
if tt.count(old)!=1:raise RuntimeError('layout test anchor not found')
test.write_text(tt.replace(old,new,1),encoding='utf-8')

new_test='''import json\nfrom pathlib import Path\nfrom types import SimpleNamespace\nfrom docx import Document\nfrom docx.shared import Pt\nfrom app.batch_prepare import _base_resume_payload\nfrom app.reference_resume_formatter import render_llm_resume\nfrom app.llm_resume_writer import SYSTEM_PROMPT, build_prompt\n\nROOT=Path(__file__).resolve().parents[1]\n\ndef _profile():\n    return json.loads((ROOT/"data/candidate_profile.json").read_text(encoding="utf-8"))\n\ndef test_uploaded_master_is_the_profile_baseline():\n    p=_profile(); exp={x["company"]:x for x in p["experience"]}\n    assert p["master_resume_reference"]["effective_date"]=="2026-10-03"\n    assert p["master_resume_reference"]["bullet_counts"]=={"Fidelity Investments":10,"Cigna Healthcare":8,"Target Corporation":8}\n    assert [len(exp[c]["evidence"]) for c in ("Fidelity Investments","Cigna Healthcare","Target Corporation")]==[10,8,8]\n    assert any("dbt transformation models" in x for x in exp["Fidelity Investments"]["evidence"])\n    assert any("AI/ML-driven analytics" in x for x in exp["Fidelity Investments"]["evidence"])\n    assert any("~120–180GB/day" in x for x in exp["Target Corporation"]["evidence"])\n    assert len(p["summary_source"])==2\n    assert p["skill_categories"]["Programming Languages"]==["Python","SQL","Scala"]\n\ndef test_base_payload_preserves_master_paragraphs_and_emphasis():\n    p=_profile(); payload=_base_resume_payload(p)\n    assert "\\n\\n" in payload["summary"]\n    assert payload["summary_emphasis"]==p["summary_emphasis"]\n    assert len(payload["experience"][0]["bullet_emphasis"])==10\n\ndef test_formatter_matches_master_visual_contract(tmp_path):\n    p=_profile(); payload=_base_resume_payload(p)\n    path=render_llm_resume(SimpleNamespace(company="Example",title="Senior Data Engineer"),p,payload,output_dir=str(tmp_path))\n    doc=Document(path); paras=doc.paragraphs\n    assert doc.styles["Normal"].font.name=="Calibri"\n    assert doc.styles["Normal"].font.size==Pt(10)\n    assert paras[0].runs[0].font.size==Pt(18)\n    assert paras[1].runs[0].font.size==Pt(13)\n    summary_heading=next(x for x in paras if x.text=="PROFESSIONAL SUMMARY")\n    assert summary_heading.runs[0].font.size==Pt(12)\n    assert summary_heading.runs[0].font.color.rgb is not None\n    summary_idx=paras.index(summary_heading)\n    summary_paras=[]\n    for x in paras[summary_idx+1:]:\n        if x.text=="TECHNICAL SKILLS":break\n        if x.text.strip():summary_paras.append(x)\n    assert len(summary_paras)==2\n    assert any(r.bold and "financial services" in r.text for x in summary_paras for r in x.runs)\n    assert sum(x.text=="Roles & Responsibilities:" for x in paras)==3\n    assert sum(x.text.startswith("Environment: ") for x in paras)==3\n    current=None;counts={"Fidelity Investments":0,"Cigna Healthcare":0,"Target Corporation":0}\n    fidelity_bullets=[]\n    for x in paras:\n        if x.text.startswith("Fidelity Investments"):current="Fidelity Investments"\n        elif x.text.startswith("Cigna Healthcare"):current="Cigna Healthcare"\n        elif x.text.startswith("Target Corporation"):current="Target Corporation"\n        elif current and x.style and "List Bullet" in x.style.name:\n            counts[current]+=1\n            if current=="Fidelity Investments":fidelity_bullets.append(x)\n    assert counts=={"Fidelity Investments":10,"Cigna Healthcare":8,"Target Corporation":8}\n    assert fidelity_bullets[8].paragraph_format.page_break_before is True\n    assert any(r.bold and "~2.5–3M" in r.text for r in fidelity_bullets[0].runs)\n    education=next(x for x in paras if x.text=="EDUCATION")\n    assert education.runs[0].font.size==Pt(12)\n\ndef test_llm_contract_uses_master_density_and_emphasis_schema():\n    assert "10 Fidelity bullets, 8 Cigna bullets, and 8 Target bullets" in SYSTEM_PROMPT\n    assert "VISUAL EMPHASIS CONTRACT" in SYSTEM_PROMPT\n    p=_profile();job=SimpleNamespace(company="Example",title="Data Engineer",description="Python SQL Databricks")\n    prompt=build_prompt(job,p,coverage_plan={})\n    schema=prompt["output_schema"]\n    assert "summary_emphasis" in schema\n    assert "bullet_emphasis" in schema["experience"][0]\n    assert prompt["quality_rules"]["fidelity_bullets"]==10\n    assert prompt["quality_rules"]["target_metric_bullets_max"]==1\n'''
(ROOT/'tests/test_master_resume_20261003.py').write_text(new_test,encoding='utf-8')

print('Applied uploaded master resume baseline and formatting contract.')
