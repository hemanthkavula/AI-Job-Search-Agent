from __future__ import annotations

import base64
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FORMAT_RECORD = ROOT / "data" / "master_word_format.json"
MASTER_RECORD = ROOT / "data" / "master_resume.json"
TEMPLATE_PATH = ROOT / "data" / "Hemanth_Kavula_Senior_Data_Engineer_Resume.docx"
CHUNK_DIR = ROOT / "data" / "master_word_template_b64"


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _style_run(run, size, *, bold=False, color=None):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_emphasized(paragraph, text, phrases, size):
    phrases = [str(x) for x in phrases or [] if str(x) and str(x) in text]
    if not phrases:
        _style_run(paragraph.add_run(text), size)
        return
    cursor = 0
    ranges = []
    for phrase in phrases:
        start = text.find(phrase)
        if start >= 0:
            ranges.append((start, start + len(phrase)))
    ranges.sort()
    merged = []
    for start, end in ranges:
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((start, end))
    for start, end in merged:
        if start > cursor:
            _style_run(paragraph.add_run(text[cursor:start]), size)
        _style_run(paragraph.add_run(text[start:end]), size, bold=True)
        cursor = end
    if cursor < len(text):
        _style_run(paragraph.add_run(text[cursor:]), size)


def _build_template_from_master() -> Path:
    master = json.loads(MASTER_RECORD.read_text(encoding="utf-8"))
    fmt = json.loads(FORMAT_RECORD.read_text(encoding="utf-8"))
    style = fmt.get("style") or {}

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(float(style.get("top_margin_in", 0.5)))
    section.bottom_margin = Inches(float(style.get("bottom_margin_in", 0.5)))
    section.left_margin = Inches(float(style.get("left_margin_in", 0.5)))
    section.right_margin = Inches(float(style.get("right_margin_in", 0.5)))

    identity = master["identity"]

    p = doc.add_paragraph()
    _style_run(p.add_run(identity["name"]), float(style.get("name_pt", 18)), bold=True, color=style.get("dark_blue_hex", "1F4E78"))

    p = doc.add_paragraph()
    _style_run(p.add_run(identity["headline"]), float(style.get("headline_pt", 13)), bold=True, color=style.get("job_blue_hex", "2E75B5"))

    contact = identity.get("contact") or {}
    p = doc.add_paragraph()
    contact_text = " | ".join(x for x in (contact.get("phone"), contact.get("email"), contact.get("linkedin")) if x)
    _style_run(p.add_run(contact_text), float(style.get("contact_pt", 11)))

    def heading(text):
        p = doc.add_paragraph()
        _style_run(p.add_run(text), float(style.get("section_heading_pt", 12)), bold=True, color=style.get("dark_blue_hex", "1F4E78"))
        return p

    heading("PROFESSIONAL SUMMARY")
    summary_emphasis = master.get("summary_emphasis") or []
    for text in master.get("summary") or []:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _add_emphasized(p, text, summary_emphasis, float(style.get("summary_pt", 11)))

    heading("TECHNICAL SKILLS")
    for category, values in (master.get("skills") or {}).items():
        p = doc.add_paragraph()
        _style_run(p.add_run(str(category)), float(style.get("body_pt", 10)), bold=True)
        _style_run(p.add_run(": " + ", ".join(str(x) for x in values)), float(style.get("body_pt", 10)))

    heading("PROFESSIONAL EXPERIENCE")
    for row in master.get("experience") or []:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(float(style.get("right_tab_in", 7.45))))
        _style_run(p.add_run(f"{row['company']} | {row['location']}"), float(style.get("company_pt", 12)), bold=True, color=style.get("dark_blue_hex", "1F4E78"))
        p.add_run("\t")
        _style_run(p.add_run(str(row.get("dates") or "")), float(style.get("body_pt", 10)), color=style.get("date_gray_hex", "666666"))

        p = doc.add_paragraph()
        _style_run(p.add_run(str(row.get("title") or "")), float(style.get("job_title_pt", 11)), bold=True, color=style.get("job_blue_hex", "2E75B5"))

        emphasis_rows = row.get("bullet_emphasis") or []
        for idx, bullet in enumerate(row.get("bullets") or []):
            p = doc.add_paragraph(style="List Paragraph")
            p.paragraph_format.left_indent = Inches(float(style.get("bullet_left_indent_in", 0.5)))
            p.paragraph_format.first_line_indent = Inches(float(style.get("bullet_hanging_indent_in", -0.25)))
            phrases = emphasis_rows[idx] if idx < len(emphasis_rows) else []
            _add_emphasized(p, bullet, phrases, float(style.get("body_pt", 10)))

        p = doc.add_paragraph()
        _style_run(p.add_run("Environment"), float(style.get("body_pt", 10)), bold=True)
        _style_run(p.add_run(": " + str(row.get("environment") or "")), float(style.get("body_pt", 10)))

    heading("EDUCATION")
    for row in master.get("education") or []:
        p = doc.add_paragraph()
        _style_run(p.add_run(str(row.get("degree") or "")), float(style.get("education_degree_pt", 11)), bold=True)
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(float(style.get("right_tab_in", 7.45))))
        school = f"{row.get('school', '')} | {row.get('location', '')}".strip(" |")
        dates = f"{row.get('start', '')} – {row.get('end', '')}".strip(" –")
        _style_run(p.add_run(school), float(style.get("body_pt", 10)))
        p.add_run("\t")
        _style_run(p.add_run(dates), float(style.get("body_pt", 10)), color=style.get("date_gray_hex", "666666"))

    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TEMPLATE_PATH)
    return TEMPLATE_PATH


def ensure_master_word_template() -> Path:
    """Materialize the master Word template.

    Prefer the exact uploaded DOCX transport when its chunks decode and match the
    recorded checksum. If that transport is unavailable or stale, rebuild a
    deterministic working template from the user-authoritative master content and
    formatting records so CI and production do not fail globally.
    """
    record = json.loads(FORMAT_RECORD.read_text(encoding="utf-8"))
    expected = (record.get("source") or {}).get("sha256")

    if TEMPLATE_PATH.exists():
        current = TEMPLATE_PATH.read_bytes()
        if expected and _sha256_bytes(current) == expected:
            return TEMPLATE_PATH

    parts = sorted(CHUNK_DIR.glob("part*.txt"))
    if parts and expected:
        try:
            encoded = "".join(p.read_text(encoding="ascii").strip() for p in parts)
            data = base64.b64decode(encoded, validate=True)
            if _sha256_bytes(data) == expected:
                TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
                TEMPLATE_PATH.write_bytes(data)
                return TEMPLATE_PATH
        except Exception:
            pass

    return _build_template_from_master()
