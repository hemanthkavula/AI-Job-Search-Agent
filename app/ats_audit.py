from __future__ import annotations
import re
from collections import Counter
from docx import Document
from app.resume_generator import jd_keywords,inferable_terms,jd_skill_terms
from app.jd_coverage_plan import build_coverage_plan

ATS_TARGET=95
HUMAN_QUALITY_TARGET=90
EXPECTED_COUNTS={"Fidelity Investments":8,"Cigna Healthcare":7,"Target Corporation":6}
METRIC_LIMITS={"Fidelity Investments":2,"Cigna Healthcare":2,"Target Corporation":0}
APPROVED_METRIC_PATTERNS={"Fidelity Investments":[r"2\.5\s*[–-]\s*3\s*(?:m|million)",r"400\s*[–-]\s*500\s*gb",r"30\s*%"],"Cigna Healthcare":[r"2\s*[–-]\s*3\s*(?:m|million)",r"25\s*[–-]\s*30\s*%"],"Target Corporation":[]}
TECH_NUMBER_PATTERNS=[r"\bs3\b",r"\bec2\b",r"\bscd\s*(?:type\s*)?2\b",r"\bgen\s*2\b",r"\badls\s*gen2\b",r"\bpython\s*3\b"]
METRIC_TOKEN_PATTERNS=[r"\b\d+(?:\.\d+)?\s*%",r"\b\d+(?:\.\d+)?\s*(?:k|m|b|million|billion)\b",r"\b\d+(?:\.\d+)?\s*(?:gb|tb|pb|mb)\b",r"\b\d+(?:\.\d+)?\s*(?:ms|milliseconds?|seconds?|minutes?|hours?)\b",r"\b(?:sub|under|below|less than)\s*[- ]?(?:second|minute|hour)\b",r"\b\d+(?:\.\d+)?\+?\s*(?:events?|records?|rows?|transactions?|members?|pipelines?|tables?|datasets?|jobs?|rules?)\b"]
TERM_ALIASES={
 "Event Hub":["Event Hub","Event Hubs","Azure Event Hub","Azure Event Hubs"],
 "Apache Kafka":["Apache Kafka","Kafka"],"Apache Spark":["Apache Spark","Spark","PySpark"],
 "Airflow":["Airflow","Apache Airflow"],"CI/CD Best Practices":["CI/CD","continuous integration","continuous delivery","continuous deployment"],
 "Batch Processing":["batch processing","batch pipeline","batch pipelines","batch workflow","batch workflows"],
 "Real-Time Data Processing":["real-time data processing","real time data processing","real-time pipeline","real-time pipelines","real time pipeline","streaming pipeline","streaming pipelines","event-driven pipeline","event-driven pipelines"],
 "Data Lineage":["Data Lineage","lineage"],"Data Quality":["Data Quality","data-quality","Great Expectations"],
 "Kinesis":["Kinesis","Amazon Kinesis"],"BigQuery":["BigQuery","Google BigQuery"],"Azure Synapse Analytics":["Azure Synapse Analytics","Azure Synapse","Synapse Analytics"],"Azure Data Factory":["Azure Data Factory","Data Factory","ADF"],"Performance Tuning":["Performance Tuning","Performance Optimization","optimized","optimizing"],"Databricks":["Databricks"],"Snowflake":["Snowflake"],"Dagster":["Dagster"],"Kubernetes":["Kubernetes"],"Terraform":["Terraform"],"CI/CD":["CI/CD","continuous integration","continuous delivery","continuous deployment"],
 "Requirements Gathering":["requirements gathering","gathering requirements","gather requirements","business requirements","technical requirements","requirements analysis"],
 "Solution Design & Development":["solution design","solution development","design and development","design, development","design/develop","designed and developed","developed and designed"],
 "Solution Implementation & Support":["implementation and support","implement and support","implemented and supported","production support","application support","solution support"],
 "Data Integration Solutions":["data integration solutions","data integrations","integration solutions","data integration","integrations"],
 "Reusable Enterprise Solutions":["reusable enterprise solutions","reusable solutions","reusable data interfaces","reusable interfaces","reusable applications","reusable apps","repurposed"],
 "Proof of Concepts":["proof of concept","proof-of-concept","proof of concepts","proof-of-concepts","poc","prototype"],
 "Customer-Facing Solutions":["customer-facing solutions","customer facing solutions","customer-facing applications","customer facing applications","customer service delivery"],
 "Application Lifecycle Management":["application lifecycle management","enterprise alm","alm process","alm practices","build environments"],
 "Enterprise ALM":["enterprise alm","application lifecycle management","alm process","alm practices"],
 "Enterprise Data Interfaces":["enterprise data interfaces","data interfaces","enterprise interfaces"],
 "Cross-Functional Collaboration":["cross-functional","cross functional","collaborate with","partner with","liaise with","stakeholders"],
 "Continuous Improvement":["continuous improvement","recommend enhancements","process improvements","platform improvements","technology roadmap"],
 "Data Modeling":["data modeling","data modelling","data model","dimensional modeling","star schema"],
 "Testing & Quality Assurance":["unit testing","integration testing","performance testing","quality assurance","qa process","testing"],
 "Security & Governance":["security and governance","secure data","security policies","access control","data governance","compliance","regulated data"]}
# Languages commonly presented as alternatives in DE JDs. A resume does not need every alternative.
OPTIONAL_LANGUAGE_ALTERNATIVES={"Go","Rust","Scala","Java"}

def document_text(path):d=Document(path);return "\n".join(p.text for p in d.paragraphs)
def _norm(s):return re.sub(r"\s+"," ",(s or "").lower()).strip()
def _literal_contains(text,term):return bool(re.search(r"(?<![a-z0-9])"+re.escape(_norm(term))+r"(?![a-z0-9])",_norm(text)))
def _contains(text,term):return any(_literal_contains(text,a) for a in TERM_ALIASES.get(term,[term]))
def _metric_sanitized(text):
 s=_norm(text)
 for pat in TECH_NUMBER_PATTERNS:s=re.sub(pat," ",s,flags=re.I)
 return s
def _is_metric_bullet(text):return any(re.search(p,_metric_sanitized(text),re.I) for p in METRIC_TOKEN_PATTERNS)
def _has_approved_metric(company,text):return any(re.search(p,_metric_sanitized(text),re.I) for p in APPROVED_METRIC_PATTERNS.get(company,[]))
def _unapproved_metric_claims(by_company):
 findings=[]
 for company,bullets in by_company.items():
  for bullet in bullets:
   s=_metric_sanitized(bullet);tokens=[m.group(0) for p in METRIC_TOKEN_PATTERNS for m in re.finditer(p,s,re.I)]
   if tokens and not _has_approved_metric(company,bullet):findings.append({"company":company,"metrics":list(dict.fromkeys(tokens)),"bullet":bullet})
   elif tokens:
    cleaned=s
    for p in APPROVED_METRIC_PATTERNS.get(company,[]):cleaned=re.sub(p," ",cleaned,flags=re.I)
    extras=[m.group(0) for p in METRIC_TOKEN_PATTERNS for m in re.finditer(p,cleaned,re.I)]
    if extras:findings.append({"company":company,"metrics":list(dict.fromkeys(extras)),"bullet":bullet})
 return findings
def _experience_bullets(paras):
 by={k:[] for k in EXPECTED_COUNTS};current=None
 for p in paras:
  t=p.text.strip()
  if t.startswith("Fidelity Investments"):current="Fidelity Investments"
  elif t.startswith("Cigna Healthcare"):current="Cigna Healthcare"
  elif t.startswith("Target Corporation"):current="Target Corporation"
  elif current and p.style and "List Bullet" in p.style.name:by[current].append(t)
 return by
def _repetition_findings(bullets):
 normalized=[re.findall(r"[a-z0-9+#.-]+",_norm(x)) for x in bullets];phrases=Counter()
 for words in normalized:
  seen=set()
  for n in (4,5,6):
   for i in range(max(0,len(words)-n+1)):
    phrase=" ".join(words[i:i+n])
    if phrase not in seen:phrases[phrase]+=1;seen.add(phrase)
 repeated=[p for p,c in phrases.items() if c>=3];openings=Counter(w[0] for w in normalized if w)
 return repeated[:10],{k:v for k,v in openings.items() if v>=5}
def _technical_skills_text(paras):
 collecting=False;rows=[]
 for p in paras:
  t=p.text.strip()
  if t.upper()=="TECHNICAL SKILLS":collecting=True;continue
  if collecting and t.upper()=="PROFESSIONAL EXPERIENCE":break
  if collecting and t:rows.append(t)
 return "\n".join(rows)
def _evidence_coverage(jd_terms,skills_text,bullets):
 exp="\n".join(bullets);listed=[t for t in jd_terms if _contains(skills_text,t)];evidenced=[t for t in listed if _contains(exp,t)];gaps=[t for t in listed if t not in evidenced]
 return listed,evidenced,gaps,round(100*len(evidenced)/max(1,len(listed)))
def _readability_score(bullets,repetition_score):
 if not bullets:return 0
 lengths=[len(re.findall(r"\b\w+[+#.-]*\b",b)) for b in bullets];score=100-min(25,sum(n>38 for n in lengths)*3)-min(20,sum(n>50 for n in lengths)*5)
 if sum(lengths)/len(lengths)>34:score-=10
 return max(40,round(min(score,repetition_score)))
def _canonical_term(term):
 aliases={
  "Synapse Analytics":"Azure Synapse Analytics",
  "Azure Synapse":"Azure Synapse Analytics",
  "Data Factory":"Azure Data Factory",
  "ETL":"ETL/ELT",
  "ELT":"ETL/ELT",
  "CI/CD":"CI/CD Best Practices",
 }
 return aliases.get(term,term)

def _dedupe_target_terms(terms):
 out=[]
 for term in terms:
  canonical=_canonical_term(term)
  if canonical not in out:out.append(canonical)
 return out

def _required_target_terms(targeted,text):
 # Alternative DE languages are not mandatory when Python already satisfies
 # the programming-language requirement.
 python_present=_contains(text,"Python")
 deduped=_dedupe_target_terms(targeted)
 return [t for t in deduped if not (python_present and t in OPTIONAL_LANGUAGE_ALTERNATIVES)]

def ats_audit(job,profile,resume_path):
 text=document_text(resume_path);low=_norm(text);plan=build_coverage_plan(job,profile);must_cover_terms=plan.get("must_cover_terms",[]);preferred_terms=plan.get("preferred_terms",[]);alternative_terms=plan.get("alternative_terms",[]);targeted=_required_target_terms(must_cover_terms,low);present=[k for k in targeted if _contains(low,k)];missing=[k for k in targeted if not _contains(low,k)];keyword_coverage=100*len(present)/max(1,len(targeted));optional_present=[k for k in preferred_terms+alternative_terms if _contains(low,k)];optional_total=len(preferred_terms)+len(alternative_terms);optional_coverage=100*len(optional_present)/max(1,optional_total) if optional_total else 100
 title_tokens=[x for x in re.findall(r"[a-z]+",_norm(job.title)) if x not in {"senior","lead","ii","iii"}];title_alignment=100 if all(x in low for x in title_tokens) else 70;sections={"professional summary","technical skills","professional experience","education"};section_score=100*sum(x in low for x in sections)/len(sections)
 paras=Document(resume_path).paragraphs;by_company=_experience_bullets(paras);bullets=[b for xs in by_company.values() for b in xs];experience_text="\n".join(bullets);plan=build_coverage_plan(job,profile);must_cover_terms=plan.get("must_cover_terms",[]);experience_covered=[t for t in must_cover_terms if _contains(experience_text,t)];experience_gaps=[t for t in must_cover_terms if not _contains(experience_text,t)];experience_coverage=100*len(experience_covered)/max(1,len(must_cover_terms));counts={k:len(v) for k,v in by_company.items()};bullet_count_score=100 if counts==EXPECTED_COUNTS else 60;metric_counts={c:sum(_is_metric_bullet(b) for b in xs) for c,xs in by_company.items()};metric_lines=sum(metric_counts.values());metric_violations={c:{"count":metric_counts[c],"limit":limit} for c,limit in METRIC_LIMITS.items() if metric_counts[c]>limit};unapproved_metrics=_unapproved_metric_claims(by_company);accomplishment_score=min(100,60+min(metric_lines,4)*10)
 if metric_violations:accomplishment_score=max(40,accomplishment_score-20*len(metric_violations))
 repeated_phrases,repeated_openings=_repetition_findings(bullets);repetition_score=100-min(35,len(repeated_phrases)*7)-min(20,sum(v-4 for v in repeated_openings.values())*4);repetition_score=max(40,repetition_score);skills_text=_technical_skills_text(paras);readability_score=_readability_score(bullets,repetition_score);human_quality_score=round(readability_score*.40+repetition_score*.25+accomplishment_score*.20+bullet_count_score*.15);compatibility_score=round(keyword_coverage*.60+title_alignment*.15+section_score*.10+bullet_count_score*.10+5);recruiter_fit_score=round(keyword_coverage*.30+experience_coverage*.35+human_quality_score*.20+title_alignment*.10+section_score*.05);score=round(compatibility_score*.70+human_quality_score*.30)
 discovery_score=getattr(job,"discovery_score",None);hands_on_cues=any(x in _norm(job.description) for x in ("hands-on","hands on","configure","configuration","implement","implementation","architecture","administration","administer","mentor","mentoring","leadership","technical leadership","operational ownership","manage platform","managing platform"));experience_gate=(not hands_on_cues) or experience_coverage>=70;gates={"discovery":discovery_score is None or discovery_score>=80,"structure":counts==EXPECTED_COUNTS,"metrics":not metric_violations and not unapproved_metrics,"repetition":repetition_score>=85,"human_quality":human_quality_score>=HUMAN_QUALITY_TARGET,"experience_depth":experience_gate};quality_gate=all(gates.values());passed=score>=ATS_TARGET and keyword_coverage>=95 and not missing and quality_gate
 return {"passed":passed,"internal_ats_score":score,"target":ATS_TARGET,"compatibility_score":compatibility_score,"recruiter_fit_score":recruiter_fit_score,"human_quality_score":human_quality_score,"human_quality_target":HUMAN_QUALITY_TARGET,"readability_score":readability_score,"keyword_coverage":round(keyword_coverage),"title_alignment":round(title_alignment),"section_score":round(section_score),"accomplishment_score":round(accomplishment_score),"targeted_jd_terms":targeted,"preferred_jd_terms":preferred_terms,"alternative_jd_terms":alternative_terms,"optional_jd_coverage":round(optional_coverage),"optional_jd_terms_present":optional_present,"missing_jd_keywords":missing,"must_cover_experience_terms":must_cover_terms,"experience_covered_terms":experience_covered,"experience_depth_gaps":experience_gaps,"experience_depth_coverage":round(experience_coverage),"metric_bearing_bullets":metric_lines,"metric_counts_by_employer":metric_counts,"metric_violations":metric_violations,"unapproved_metric_claims":unapproved_metrics,"approved_metric_patterns":APPROVED_METRIC_PATTERNS,"bullet_counts":counts,"bullet_count_score":bullet_count_score,"skills_taxonomy_score":100,"repetition_score":repetition_score,"repeated_phrases":repeated_phrases,"repeated_opening_verbs":repeated_openings,"discovery_score":discovery_score,"quality_gate_passed":quality_gate,"quality_gates":gates,"status":"ATS_PASS" if passed else "HOLD_ATS_REVIEW","note":"Internal estimates only. internal_ats_score measures textual/structural ATS alignment; recruiter_fit_score additionally weights demonstrated must-cover requirements in Professional Experience. Neither represents a proprietary employer ATS score."}
