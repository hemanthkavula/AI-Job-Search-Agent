from __future__ import annotations

from pathlib import Path
import os
import re
import shutil
import subprocess
import tempfile
import time
from docx import Document


def _find_office() -> str | None:
    """Find a LibreOffice/soffice executable without launching Microsoft Word."""
    candidates = [
        shutil.which("libreoffice"),
        shutil.which("soffice"),
        r"C:\\Program Files\\LibreOffice\\program\\soffice.exe",
        r"C:\\Program Files (x86)\\LibreOffice\\program\\soffice.exe",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return str(candidate)
    return None


def _native_convert(src: Path, target: Path) -> tuple[bool, str]:
    """Render one approved DOCX headlessly using an isolated LibreOffice profile."""
    office = _find_office()
    if not office:
        return False, (
            "LibreOffice/soffice was not found. Install LibreOffice for unattended "
            "DOCX-to-PDF conversion; Microsoft Word COM is intentionally disabled."
        )

    target.parent.mkdir(parents=True, exist_ok=True)
    # Keep the LibreOffice user profile OUTSIDE the resume output directory.
    # OneDrive-synced job folders can delay/lock profile files on Windows and cause
    # soffice to exit successfully without leaving the requested PDF.
    profile_dir = Path(tempfile.mkdtemp(prefix="ai_job_resume_lo_"))
    profile_uri = profile_dir.resolve().as_uri()
    cmd = [
        office,
        "--headless",
        "--nologo",
        "--nodefault",
        "--nolockcheck",
        "--norestore",
        f"-env:UserInstallation={profile_uri}",
        "--convert-to",
        "pdf:writer_pdf_Export",
        "--outdir",
        str(target.parent),
        str(src),
    ]
    env = dict(os.environ)
    env["SAL_DISABLE_SYNCHRONOUS_PRINTER_DETECTION"] = "1"
    try:
        proc = subprocess.run(cmd, check=False, timeout=90, capture_output=True, text=True, env=env)
        detail = "\n".join(x.strip() for x in (proc.stdout, proc.stderr) if x and x.strip())
        if proc.returncode != 0:
            return False, f"LibreOffice conversion failed (exit {proc.returncode}): {detail or 'no diagnostic output'}"

        # On Windows/OneDrive, filesystem visibility can lag behind soffice exit.
        for _ in range(20):
            if target.exists() and target.stat().st_size > 0:
                return True, "ok"
            time.sleep(0.25)
        return False, (
            "LibreOffice exited successfully but PDF was not created"
            + (f": {detail}" if detail else "")
        )
    except subprocess.TimeoutExpired:
        return False, "LibreOffice conversion timed out after 90 seconds"
    except Exception as exc:
        return False, f"LibreOffice conversion failed to start: {exc}"
    finally:
        def _onerror(func, path, exc_info):
            try:
                os.chmod(path, 0o700)
                func(path)
            except OSError:
                pass
        shutil.rmtree(profile_dir, onerror=_onerror)


def _conversion_attempt(src: Path, target: Path) -> tuple[bool, str]:
    try:
        if target.exists():
            target.unlink()
    except OSError as exc:
        return False, f"Existing PDF could not be replaced: {exc}"
    return _native_convert(src, target)


def convert_docx_to_pdf(docx_path: str, attempts: int = 2) -> str | None:
    """Convert the same approved DOCX, retrying rendering only; never rebuild content."""
    src = Path(docx_path).resolve()
    if not src.exists() or src.stat().st_size == 0:
        print(f"PDF conversion skipped | DOCX missing or empty: {src}", flush=True)
        return None

    target = src.with_suffix(".pdf")
    last_reason = "conversion not attempted"
    for attempt in range(1, max(1, attempts) + 1):
        ok, last_reason = _conversion_attempt(src, target)
        if ok:
            if attempt > 1:
                print(f"PDF conversion recovered on attempt {attempt}", flush=True)
            return str(target)
        print(
            f"PDF conversion attempt {attempt}/{max(1, attempts)} failed | {last_reason}",
            flush=True,
        )
        if attempt < max(1, attempts):
            time.sleep(1.0)

    print(f"PDF conversion exhausted retries | {last_reason}", flush=True)
    return None

def _docx_signature(docx_path: str) -> dict:
    doc = Document(str(docx_path))
    paragraphs = [re.sub(r"\s+", " ", p.text).strip() for p in doc.paragraphs if p.text.strip()]
    bullets = sum(
        1 for p in doc.paragraphs
        if p.text.strip() and p.style and "List Bullet" in p.style.name
    )
    names = {"PROFESSIONAL SUMMARY", "TECHNICAL SKILLS", "PROFESSIONAL EXPERIENCE", "EDUCATION"}
    sections = [x.upper() for x in paragraphs if x.upper() in names]
    return {"paragraphs": paragraphs, "bullets": bullets, "sections": sections}


def _pdf_text(pdf_path: str) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(pdf_path))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        return re.sub(r"\s+", " ", text).strip(), len(reader.pages)
    except Exception:
        return "", 0


def _tokens(text: str) -> list[str]:
    # Token parity tolerates harmless PDF extraction differences in bullets,
    # punctuation, Unicode dashes, and line wrapping without weakening content checks.
    return re.findall(r"[a-z0-9+#./%-]+", text.lower())


def _paragraph_covered(paragraph: str, pdf_tokens: list[str]) -> bool:
    wanted = _tokens(paragraph)
    if not wanted:
        return True
    pdf_set = set(pdf_tokens)
    # Short headings/labels should be exact token subsets. Longer material paragraphs
    # may differ slightly in extraction while still containing the same rendered text.
    ratio = sum(1 for token in wanted if token in pdf_set) / len(wanted)
    threshold = 1.0 if len(wanted) <= 4 else 0.97
    return ratio >= threshold


def validate_docx_pdf_parity(docx_path: str, pdf_path: str | None) -> dict:
    if not pdf_path or not Path(pdf_path).exists() or Path(pdf_path).stat().st_size == 0:
        return {
            "passed": False, "reason": "PDF was not created", "text_coverage": 0,
            "sections_match": False, "page_count": 0,
        }

    sig = _docx_signature(docx_path)
    pdf_text, pages = _pdf_text(pdf_path)
    if not pdf_text:
        return {
            "passed": False, "reason": "PDF text could not be validated",
            "text_coverage": 0, "sections_match": False, "page_count": pages,
        }

    pdf_tokens = _tokens(pdf_text)
    material = [p for p in sig["paragraphs"] if len(p) >= 8]
    matched = sum(1 for p in material if _paragraph_covered(p, pdf_tokens))
    coverage = round(100 * matched / max(1, len(material)), 1)

    pdf_token_set = set(pdf_tokens)
    sections_match = all(set(_tokens(s)).issubset(pdf_token_set) for s in sig["sections"])
    passed = coverage >= 95 and sections_match and pages > 0

    return {
        "passed": passed,
        "reason": None if passed else "DOCX/PDF material text or section mismatch",
        "text_coverage": coverage,
        "sections_match": sections_match,
        "docx_bullet_count": sig["bullets"],
        "sections": sig["sections"],
        "page_count": pages,
        "renderer": "libreoffice_headless",
    }
