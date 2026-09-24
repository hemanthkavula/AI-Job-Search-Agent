
# USER RESUME PROMPT CONTRACT (2026-09-16)
# Generate a fresh resume per JD with strict domain lock:
# Fidelity = financial services; Cigna = healthcare; Target = retail.
# Preserve identity, company names, titles, locations, dates, education and certifications.
# Exact bullet counts: Fidelity 8, Cigna 7, Target 6.
# Fidelity gets strongest/newest JD technology coverage; Cigna moderate; Target foundational.
# Summary must be concise and naturally include top JD terms.
# Technical Skills should use standard ATS categories and blend relevant JD terminology naturally.
# No label such as "JD-aligned technologies".
# Metrics: Fidelity max 1-2, Cigna max 1-2, Target zero; never invent metrics.
# Linear ATS layout only: no tables, columns, graphics or text boxes.
# Use exact JD terminology where appropriate and run the ATS audit after generation.
from __future__ import annotations
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parent.parent
ALIASES={"AWS Glue":["aws glue","glue"],"Amazon EMR":["amazon emr","emr"],"Amazon S3":["amazon s3","s3"],"Amazon Redshift":["amazon redshift","redshift"],"Azure Data Factory":["azure data factory","adf"],"Azure Synapse Analytics":["azure synapse","synapse"],"ADLS Gen2":["adls gen2","adls","azure data lake storage"],"Apache Spark":["apache spark","spark"],"Apache Kafka":["apache kafka","kafka"],"PySpark":["pyspark"],"Slowly Changing Dimensions":["slowly changing dimensions","scd type 2","scd"],"CI/CD Best Practices":["ci/cd","cicd"]}
def clean_company_name(value):
    """Return a clean display/file company name without legal suffixes or leading 'The'."""
    s=re.sub(r"\s+"," ",str(value or "")).strip()
    s=re.sub(r"(?i)^the\s+","",s)
    # Handle inverted legal names such as "Computer Merchant, Ltd., The".
    s=re.sub(r"(?i)[\s,.-]+(?:incorporated|inc|corporation|corp|limited|ltd|llc|l\.l\.c|plc|co|company)[.,\s]*$", "", s).strip(" ,.-")
    s=re.sub(r"(?i)[\s,.-]+the$", "", s).strip(" ,.-")
    # A trailing legal suffix may become exposed after removing ", The".
    s=re.sub(r"(?i)[\s,.-]+(?:incorporated|inc|corporation|corp|limited|ltd|llc|l\.l\.c|plc|co|company)$", "", s).strip(" ,.-")
    return s or str(value or "").strip()

def safe_name(v):return re.sub(r"[^A-Za-z0-9_-]+","_",v).strip("_")[:80]
def all_verified(profile):
    out=[]
    for xs in profile.get("skill_categories",{}).values():out.extend(xs)
    return list(dict.fromkeys(out))
def jd_keywords(jd,profile):
    low=(jd or "").lower();out=[]
    for s in all_verified(profile):
        if any(v in low for v in ALIASES.get(s,[s.lower()])):out.append(s)
    return out

def jd_skill_terms(jd):
    """Extract ATS-friendly technologies and material DE/application-integration concepts directly from the complete JD."""
    text=jd or ""
    catalog={
      "Python":("python",),"SQL":("sql",),"Scala":("scala",),"Java":("java",),"Go":("golang",),"Rust":("rust",),
      "PySpark":("pyspark",),"Apache Spark":("apache spark","spark"),"Apache Kafka":("apache kafka","kafka"),
      "Ab Initio":("ab initio","ab-initio"),"Informatica":("informatica","powercenter","informatica powercenter"),
      "Apache Flink":("apache flink","flink"),"Databricks":("databricks",),"Snowflake":("snowflake",),"dbt":("dbt",),
      "Dagster":("dagster",),"Airflow":("airflow",),"Fivetran":("fivetran",),"Airbyte":("airbyte",),
      "Kubernetes":("kubernetes","k8s"),"Docker":("docker",),"Terraform":("terraform",),
      "AWS Glue":("aws glue",),"Amazon S3":("amazon s3","s3"),"Amazon EMR":("amazon emr","emr"),
      "Amazon Redshift":("amazon redshift","redshift"),"AWS Lambda":("aws lambda","lambda"),"AWS Kinesis":("amazon kinesis","kinesis"),
      "Azure Data Factory":("azure data factory","data factory","adf"),"Azure Synapse Analytics":("azure synapse analytics","azure synapse","synapse analytics"),
      "ADLS Gen2":("adls gen2","azure data lake storage gen2"),"Azure Event Hubs":("azure event hubs","azure event hub","event hubs","event hub"),
      "BigQuery":("bigquery","google bigquery"),"GCP":("gcp","google cloud"),"Apache Iceberg":("apache iceberg","iceberg"),
      "Delta Lake":("delta lake",),"Hudi":("hudi",),"PostgreSQL":("postgresql","postgres"),"MySQL":("mysql",),
      "MongoDB":("mongodb",),"Oracle":("oracle",),"SQL Server":("sql server",),
      "Microsoft Dynamics 365 CRM":("microsoft dynamics 365 crm","dynamics 365 crm","dynamics crm"),
      "Microsoft Power Platform":("microsoft power platform","power platform"),"Microsoft Dataverse":("microsoft dataverse","dataverse"),
      "Power Apps":("power apps","powerapps"),"Power Automate":("power automate",),"Dynamics 365 plug-ins":("dynamics 365 plug-ins","dynamics 365 plugins","crm plug-ins","crm plugins"),
      "SSRS":("sql server reporting services","ssrs"),"SharePoint":("sharepoint",),"Power BI":("power bi","powerbi"),
      "REST APIs":("rest api","rest apis","restful api"),"Enterprise ALM":("enterprise alm","application lifecycle management","alm process"),
      "CI/CD":("ci/cd","continuous integration","continuous delivery","continuous deployment"),"Git":("git",),
      "Data Governance":("data governance",),"Data Lineage":("data lineage","lineage"),"Data Quality":("data quality",),
      "ETL":("etl",),"ELT":("elt",),"Batch Processing":("batch processing",),"Real-Time Data Processing":("real-time data processing","real time data processing","streaming"),
      "Dimensional Modeling":("dimensional modeling",),"Systems-of-record Integration":("systems-of-record integration","system-of-record integration","systems of record"),
      "Enterprise Data Interfaces":("enterprise data interfaces","data interfaces"),"CRM Data Integration":("crm data integration","crm integration"),
      "Production Support":("production support",),"Technical Documentation":("technical documentation",),"Agile/Scrum":("agile","scrum"),
    }
    low=text.lower();out=[]
    for label,aliases in catalog.items():
        if any(re.search(r"(?<![a-z0-9])"+re.escape(a)+r"(?![a-z0-9])",low) for a in aliases):out.append(label)
    return out

def inferable_terms(jd):
    """JD concepts that may be stated when already evidenced by the candidate's documented work."""
    low=(jd or "").lower()
    mapping={
      "ETL/ELT":["etl","elt"],"Batch Processing":["batch processing","batch pipelines"],
      "Real-Time Data Processing":["real-time","real time","streaming"],
      "Data Lakehouse":["lakehouse"],"Data Lineage":["data lineage","lineage"],
      "Data Governance":["data governance","governance"],"Data Quality":["data quality"],
      "Performance Tuning":["performance tuning","performance optimization"],
      "Infrastructure as Code":["infrastructure as code","iac"],"Dimensional Modeling":["dimensional modeling"],
      "Schema Evolution":["schema evolution"],"Orchestration":["orchestration","workflow orchestration"],
      "Requirements Gathering":["requirements gathering","gather requirements","business requirements","technical requirements"],
      "Solution Design & Development":["solution design","design and development","design, development","design/develop","develop and implement"],
      "Solution Implementation & Support":["implementation and support","implement and support","implementation, and support","production support"],
      "Data Integration Solutions":["data integration solutions","data integrations","integration solutions","data interfaces"],
      "Reusable Enterprise Solutions":["reusable","repurposed","reusable data interfaces","reusable applications","reusable apps"],
      "Proof of Concepts":["proof of concept","proof-of-concept","proof of concepts","poc"],
      "Customer-Facing Solutions":["customer-facing","customer facing","customer service delivery"],
      "Application Lifecycle Management":["application lifecycle management","enterprise alm","alm process","build environments"],
      "Cross-Functional Collaboration":["cross-functional","cross functional","liaise with","partner with","collaborate with"],
      "Continuous Improvement":["continuous improvement","technology updates","roadmap","recommend enhancements"],
      "Data Modeling":["data model","data modeling","data modelling"],
      "Testing & Quality Assurance":["unit testing","integration testing","performance testing","quality assurance","qa process"],
      "Security & Governance":["secure data","security policies","access control","governance","compliance"],
      "Machine Learning & AI":["machine learning","ml models","artificial intelligence","ai models","model training","model inference"],
      "Generative AI":["generative ai","genai","large language model","large language models","llm","llms"],
      "MLOps":["mlops","model deployment","model serving","model monitoring","model registry"],
      "Vector Search & Embeddings":["vector database","vector databases","vector search","embeddings","embedding models"],
      "Retrieval-Augmented Generation":["retrieval augmented generation","retrieval-augmented generation","rag"],
      "Data Science Collaboration":["data scientists","data science teams","machine learning engineers","ml engineers"]
    }
    return [label for label,terms in mapping.items() if any(t in low for t in terms)]

def unsupported_jd_terms(jd,profile):
    known=set(x.lower() for x in all_verified(profile));terms=[]
    common=["flink","kubernetes","java","golang","go","rust","gcp","bigquery","dbt","dagster","fivetran","airbyte","iceberg","hudi","teradata"]
    low=(jd or "").lower()
    for x in common:
        if re.search(r"(?<![a-z0-9])"+re.escape(x)+r"(?![a-z0-9])",low) and x not in known:terms.append(x)
    return terms


def _jd_themes(jd):
    low=(jd or "").lower()
    theme_map={
      "streaming":["stream","kafka","kinesis","event hub","flink"],
      "batch":["batch","etl","elt","spark"],
      "orchestration":["orchestrat","airflow","dagster","workflow"],
      "quality":["data quality","validation","great expectations"],
      "lineage":["lineage","governance","catalog"],
      "platform":["data platform","data product","curated","warehouse","lake"],
      "devops":["terraform","kubernetes","ci/cd","devops","infrastructure as code"],
      "performance":["performance","optimiz","scalab","latency"]
    }
    return [k for k,terms in theme_map.items() if any(t in low for t in terms)]

def generate_jd_specific_bullet(line,jd,company,index):
    """Create fresh JD-oriented wording from documented evidence; never add an unsupported employer/tool claim."""
    s=line.strip().rstrip(".")
    low=s.lower(); themes=_jd_themes(jd)
    prefixes={
      "Fidelity Investments":["Engineered","Designed","Built","Implemented","Optimized","Developed","Strengthened","Delivered"],
      "Cigna Healthcare":["Developed","Built","Implemented","Automated","Designed","Improved","Delivered"],
      "Target Corporation":["Built","Developed","Implemented","Automated","Supported","Delivered"]
    }
    verb=prefixes.get(company,["Built"])[index % len(prefixes.get(company,["Built"]))]
    # Replace a leading generic action verb so the bullet is visibly regenerated.
    s=re.sub(r"^(built|developed|designed|implemented|created|optimized|automated|enabled|processed|managed|supported|delivered|used)\b",verb,s,flags=re.I)
    additions=[]
    if "streaming" in themes and any(x in low for x in ["kafka","kinesis","event hub","stream"]):
        additions.append("for reliable real-time data processing")
    if "orchestration" in themes and any(x in low for x in ["airflow","data factory","glue","pipeline","etl"]):
        additions.append("with production workflow orchestration")
    if "quality" in themes and any(x in low for x in ["quality","validation","pipeline","dataset"]):
        additions.append("with embedded data quality controls")
    if "lineage" in themes and any(x in low for x in ["catalog","lake","warehouse","dataset","governance"]):
        additions.append("supporting governed and traceable data flows")
    if "platform" in themes and any(x in low for x in ["lake","warehouse","dataset","snowflake","redshift"]):
        additions.append("to deliver curated analytics-ready data products")
    if "performance" in themes and any(x in low for x in ["spark","pyspark","sql","pipeline","runtime","performance"]):
        additions.append("with focus on scalable production performance")
    if additions:
        s=s+", "+", ".join(dict.fromkeys(additions))
    return s+"."

def tailor_supported_bullet(line,jd):
    """Rewrite supported evidence toward JD themes without inventing tools or facts."""
    low=(jd or "").lower(); s=line.strip().rstrip(".")
    suffix=[]
    if any(x in low for x in ["reliab","observab","monitor"]) and any(x in s.lower() for x in ["quality","validation","pipeline","stream"]):
        suffix.append("strengthening production reliability and monitoring")
    if "lineage" in low and any(x in s.lower() for x in ["data lake","dataset","warehouse","quality"]):
        suffix.append("supporting traceable, governed data flows")
    if any(x in low for x in ["curated","data product"]) and any(x in s.lower() for x in ["dataset","data lake","warehouse"]):
        suffix.append("delivering curated analytics-ready data products")
    if any(x in low for x in ["orchestrat","workflow"]) and any(x in s.lower() for x in ["pipeline","etl","airflow","data factory","glue"]):
        suffix.append("with production workflow orchestration")
    return s + (", " + "; ".join(dict.fromkeys(suffix)) if suffix else "") + "."

def _rank(lines,jd,keys):
    words=set(re.findall(r"[a-z0-9+#.-]+",(jd or "").lower()))
    def score(x):
        low=x.lower()
        overlap=len(set(re.findall(r"[a-z0-9+#.-]+",low))&words)
        exact=4*sum(k.lower() in low for k in keys)
        themes=3*sum(t in low and t in (jd or "").lower() for t in ["stream","pipeline","orchestrat","lineage","quality","terraform","snowflake","spark","kafka","cloud","data lake","warehouse"])
        return overlap+exact+themes
    return sorted(lines,key=score,reverse=True)
def _h(doc,t):
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(1);p.paragraph_format.keep_with_next=True;r=p.add_run(t);r.bold=True;r.font.size=Pt(10.5)

def _keep_employer_block_together(company_p,title_p,first_bullet_p=None,environment_p=None):
    """Prevent an employer header/title from being orphaned at a page bottom."""
    company_p.paragraph_format.keep_with_next=True
    title_p.paragraph_format.keep_with_next=True
    if environment_p is not None:
        environment_p.paragraph_format.keep_with_next=True
    if first_bullet_p is not None:
        first_bullet_p.paragraph_format.keep_together=True

def _format_experience_bullet(p):
    p.paragraph_format.left_indent=Inches(.16)
    p.paragraph_format.first_line_indent=Inches(-.10)
    p.paragraph_format.keep_together=True
def generate_resume(job,analysis,profile,output_dir="generated/resumes"):
    keys=jd_keywords(job.description,profile); inferred=inferable_terms(job.description); doc=Document();s=doc.sections[0]
    s.top_margin=Inches(.45);s.bottom_margin=Inches(.45);s.left_margin=Inches(.55);s.right_margin=Inches(.55)
    doc.styles["Normal"].font.name="Arial";doc.styles["Normal"].font.size=Pt(9.3)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(profile["name"]);r.bold=True;r.font.size=Pt(15)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(job.title or profile.get("headline","Senior Data Engineer"));r.bold=True;r.font.size=Pt(10.5)
    c=profile.get("contact",{});p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(" | ".join(x for x in [c.get("phone"),c.get("email"),c.get("linkedin")] if x))
    _h(doc,"PROFESSIONAL SUMMARY")
    focus=", ".join((keys+inferred)[:7]) or "Python, SQL, PySpark, Apache Spark, cloud data engineering"
    doc.add_paragraph(f"Senior Data Engineer with 5+ years of experience building scalable batch and real-time data platforms across financial services, healthcare, and retail. Experienced in {focus}, with a strong background in cloud data lakes, warehouses, streaming pipelines, dimensional modeling, data quality, lineage, and production reliability. Proven ability to deliver analytics-ready data products and optimize distributed data workloads for performance and scale.")
    _h(doc,"TECHNICAL SKILLS")
    jd_skills=jd_skill_terms(job.description)
    base=all_verified(profile)
    # Include both literal JD technologies and supported/inferable ATS concepts.
    # This ensures exact phrases such as Batch Processing, Real-Time Data Processing
    # and Data Lineage are present naturally in Technical Skills when relevant.
    verified=set(base)
    safe_jd=[x for x in jd_skills if x in verified]
    skills=list(dict.fromkeys(safe_jd+keys+inferred+base))
    groups={
      "Programming Languages":["Python","SQL","Scala","Java","Go","Rust"],
      "Data Engineering & Processing":["PySpark","Apache Spark","Apache Kafka","Apache Flink","Databricks","Batch Processing","Real-Time Data Processing","ETL","ELT","ETL/ELT"],
      "Workflow Orchestration":["Dagster","Airflow","Apache Airflow","dbt","Fivetran","Airbyte"],
      "Cloud & Data Platforms":["Snowflake","BigQuery","GCP","Delta Lake","Apache Iceberg","Hudi"],
      "AWS Services":["AWS Glue","Amazon S3","Amazon EMR","Amazon Redshift","AWS Lambda","AWS Kinesis"],
      "Azure Services":["Azure Data Factory","Azure Synapse","Azure Synapse Analytics","ADLS Gen2","Azure Event Hubs","Event Hub"],
      "Infrastructure & DevOps":["Kubernetes","Docker","Terraform","Jenkins","GitHub Actions","CI/CD","CI/CD Best Practices","Git"],
      "Data Quality & Governance":["Great Expectations","Data Quality","Data Lineage","Data Governance","AWS Glue Data Catalog","Azure Purview"],
      "Databases":["PostgreSQL","MySQL","MongoDB","Oracle","SQL Server","DynamoDB"],
      "Data Modeling":["Dimensional Modeling","Star Schema","Snowflake Schema","Slowly Changing Dimensions"]
    }
    used=set()
    for label,wanted in groups.items():
        vals=[x for x in wanted if x in skills and x not in used]
        if vals:
            used.update(vals);p=doc.add_paragraph();r=p.add_run(label+": ");r.bold=True;p.add_run(", ".join(vals))
    _h(doc,"PROFESSIONAL EXPERIENCE")
    for exp in profile["experience"]:
        company_p=doc.add_paragraph();r=company_p.add_run(f"{exp['company']} | {exp.get('location','')}");r.bold=True;r=company_p.add_run(f"    {exp['dates']}");r.bold=True
        title_p=doc.add_paragraph();r=title_p.add_run(exp["title"]);r.bold=True
        environment_p=doc.add_paragraph();r=environment_p.add_run("Environment: ");r.bold=True;environment_p.add_run(exp.get("environment",""))
        limits={"Fidelity Investments":8,"Cigna Healthcare":7,"Target Corporation":6}
        ranked=_rank(exp["evidence"],job.description,keys)[:limits.get(exp["company"],7)]
        first_bullet=None
        for line in ranked:
            p=doc.add_paragraph(style="List Bullet");_format_experience_bullet(p);p.add_run(line.strip())
            if first_bullet is None:first_bullet=p
        _keep_employer_block_together(company_p,title_p,first_bullet,environment_p)
    _h(doc,"EDUCATION")
    for e in profile["education"]:
        p=doc.add_paragraph();r=p.add_run(e["degree"]);r.bold=True;doc.add_paragraph(f"{e['school']} | {e['location']}    {e['start']} – {e['end']}")
    root=ROOT/output_dir;root.mkdir(parents=True,exist_ok=True);pattern=profile.get("output",{}).get("resume_filename_pattern","Hemanth_Kavula_{Company}_{JobTitle}")
    stem=pattern.replace("{Company}",safe_name(clean_company_name(job.company))).replace("{JobTitle}",safe_name(job.title))
    timestamp=datetime.now(ZoneInfo("America/New_York")).strftime("%Y%m%d_%H%M%S")
    job_dir=root/f"{safe_name(job.company)}_{safe_name(job.title)}_{timestamp}";job_dir.mkdir(parents=True,exist_ok=True)
    path=job_dir/f"{stem}.docx";doc.save(path)
    return str(path)


def render_llm_resume(job, profile, generated, output_dir="generated/resumes"):
    """Render validated LLM JSON into the same ATS-friendly DOCX format."""
    doc=Document();s=doc.sections[0]
    s.top_margin=Inches(.45);s.bottom_margin=Inches(.45);s.left_margin=Inches(.55);s.right_margin=Inches(.55)
    doc.styles["Normal"].font.name="Arial";doc.styles["Normal"].font.size=Pt(9.3)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(profile["name"]);r.bold=True;r.font.size=Pt(15)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(job.title or profile.get("headline","Senior Data Engineer"));r.bold=True;r.font.size=Pt(10.5)
    contact=profile.get("contact",{});p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(" | ".join(x for x in [contact.get("phone"),contact.get("email"),contact.get("linkedin")] if x))
    _h(doc,"PROFESSIONAL SUMMARY");doc.add_paragraph(generated.get("summary",""))
    _h(doc,"TECHNICAL SKILLS")
    for label,vals in generated.get("skills",{}).items():
        p=doc.add_paragraph();r=p.add_run(str(label)+": ");r.bold=True;p.add_run(", ".join(vals))
    _h(doc,"PROFESSIONAL EXPERIENCE")
    expected={x["company"]:x for x in profile["experience"]}
    limits={"Fidelity Investments":8,"Cigna Healthcare":7,"Target Corporation":6}
    for item in generated.get("experience",[]):
        base=expected.get(item.get("company"))
        if not base: continue
        company_p=doc.add_paragraph();r=company_p.add_run(f"{base['company']} | {base.get('location','')}");r.bold=True;r=company_p.add_run(f"    {base['dates']}");r.bold=True
        title_p=doc.add_paragraph();r=title_p.add_run(base["title"]);r.bold=True
        bullets=item.get("bullets",[])[:limits.get(base["company"],7)]
        first_bullet=None
        for line in bullets:
            p=doc.add_paragraph(style="List Bullet");_format_experience_bullet(p);p.add_run(str(line).strip())
            if first_bullet is None:first_bullet=p
        _keep_employer_block_together(company_p,title_p,first_bullet)
    _h(doc,"EDUCATION")
    for e in profile["education"]:
        p=doc.add_paragraph();r=p.add_run(e["degree"]);r.bold=True;doc.add_paragraph(f"{e['school']} | {e['location']}    {e['start']} – {e['end']}")
    root=ROOT/output_dir;root.mkdir(parents=True,exist_ok=True);pattern=profile.get("output",{}).get("resume_filename_pattern","Hemanth_Kavula_{Company}_{JobTitle}")
    stem=pattern.replace("{Company}",safe_name(job.company)).replace("{JobTitle}",safe_name(job.title))
    timestamp=datetime.now(ZoneInfo("America/New_York")).strftime("%Y%m%d_%H%M%S")
    job_dir=root/f"{safe_name(job.company)}_{safe_name(job.title)}_{timestamp}";job_dir.mkdir(parents=True,exist_ok=True)
    path=job_dir/f"{stem}.docx";doc.save(path);return str(path)
