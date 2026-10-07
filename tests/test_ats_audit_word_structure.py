from docx import Document

from app.ats_audit import EXPECTED_COUNTS, _experience_bullets, _experience_cloud_text
from app.master_resume import load_master_resume


def _paragraphs_with_direct_formatted_experience():
    doc = Document()
    for company, count in EXPECTED_COUNTS.items():
        doc.add_paragraph(f"{company} | Location")
        doc.add_paragraph("Data Engineer")
        doc.add_paragraph("Roles & Responsibilities:")
        for index in range(count):
            # Deliberately keep the default Normal style. The authoritative
            # user-uploaded Word template uses direct paragraph formatting for
            # its visible bullets rather than the named List Bullet style.
            doc.add_paragraph(f"Built governed data pipeline responsibility {index + 1}")
        doc.add_paragraph(f"Skills: {company} tools")
    doc.add_paragraph("EDUCATION")
    return doc.paragraphs


def _paragraphs_with_current_no_label_experience():
    doc = Document()
    master = load_master_resume()
    for row in master["experience"]:
        doc.add_paragraph(f'{row["company"]} | {row["location"]}')
        doc.add_paragraph(f'{row["title"]}\t{row["dates"]}')
        for index in range(len(row["bullets"])):
            doc.add_paragraph(f"Built governed data pipeline responsibility {index + 1}")
        doc.add_paragraph(f'Environment: {row["company"]} tools')
    doc.add_paragraph("EDUCATION")
    return doc.paragraphs


def test_experience_audit_reads_current_no_roles_label_layout():
    paragraphs = _paragraphs_with_current_no_label_experience()
    by_company = _experience_bullets(paragraphs)
    assert {company: len(rows) for company, rows in by_company.items()} == EXPECTED_COUNTS


def test_experience_audit_reads_master_word_structure_without_list_bullet_style():
    paragraphs = _paragraphs_with_direct_formatted_experience()
    by_company = _experience_bullets(paragraphs)
    assert {company: len(rows) for company, rows in by_company.items()} == EXPECTED_COUNTS


def test_experience_cloud_text_includes_structural_bullets_and_skills_footer():
    paragraphs = _paragraphs_with_direct_formatted_experience()
    cloud_text = _experience_cloud_text(paragraphs)
    for company, expected_count in EXPECTED_COUNTS.items():
        assert cloud_text[company].count("responsibility") == expected_count
        assert f"Skills: {company} tools" in cloud_text[company]
