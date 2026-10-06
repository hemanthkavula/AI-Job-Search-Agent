from docx import Document
from docx.shared import Inches

from app.master_resume import load_master_resume
from app.resume_pagination import enforce_experience_start_rule


def _paragraph_by_prefix(doc, prefix):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))


def test_employer_header_title_and_roles_share_one_left_edge(tmp_path):
    path = tmp_path / "resume.docx"
    doc = Document()

    for index, row in enumerate(load_master_resume()["experience"]):
        header = doc.add_paragraph(f"{row['company']} | {row['location']}")
        header.paragraph_format.left_indent = Inches(0.10 + index * 0.05)

        title = doc.add_paragraph(f"{row['title']}\t{row['dates']}")
        title.paragraph_format.left_indent = Inches(0.06 + index * 0.03)

        roles = doc.add_paragraph("Roles & Responsibilities:")
        roles.paragraph_format.left_indent = Inches(0.04 + index * 0.02)

        first = doc.add_paragraph(f"First bullet for {row['company']}")
        first.paragraph_format.left_indent = Inches(0.50)
        doc.add_paragraph(f"Second bullet for {row['company']}")
        doc.add_paragraph("Environment: Python, SQL")

    doc.save(path)

    result = enforce_experience_start_rule(path)
    assert result["passed"] is True, result["reasons"]

    fixed = Document(path)
    for row in load_master_resume()["experience"]:
        header = _paragraph_by_prefix(fixed, row["company"])
        title = _paragraph_by_prefix(fixed, row["title"])
        roles = _paragraph_by_prefix(fixed, "Roles & Responsibilities:")

        assert header.paragraph_format.left_indent.inches == 0
        assert title.paragraph_format.left_indent.inches == 0
        assert roles.paragraph_format.left_indent.inches == 0
        assert header.paragraph_format.keep_with_next is True
        assert title.paragraph_format.keep_with_next is True
        assert roles.paragraph_format.keep_with_next is True
