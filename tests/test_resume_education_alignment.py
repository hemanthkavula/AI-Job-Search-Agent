from pathlib import Path
import shutil

from docx import Document
from docx.shared import Inches

from app.reference_resume_formatter import WORD_TEMPLATE_PATH
from app.resume_pagination import enforce_experience_start_rule, validate_experience_start_rule


def _left_inches(paragraph):
    value = paragraph.paragraph_format.left_indent
    return 0.0 if value is None else value.inches


def test_education_degree_and_university_are_normalized_to_same_left_edge(tmp_path):
    target = Path(tmp_path) / "resume.docx"
    shutil.copyfile(WORD_TEMPLATE_PATH, target)

    doc = Document(target)
    degree = next(p for p in doc.paragraphs if p.text.strip() == "Master of Science in Computer Science")
    school = next(p for p in doc.paragraphs if p.text.strip().startswith("Rowan University"))
    degree.paragraph_format.left_indent = Inches(0.20)
    school.paragraph_format.left_indent = Inches(0.35)
    doc.save(target)

    result = enforce_experience_start_rule(target)
    assert result["passed"], result["reasons"]

    validated = validate_experience_start_rule(target)
    assert validated["passed"], validated["reasons"]
    assert "same left edge" in validated["education_policy"]

    doc = Document(target)
    degree = next(p for p in doc.paragraphs if p.text.strip() == "Master of Science in Computer Science")
    school = next(p for p in doc.paragraphs if p.text.strip().startswith("Rowan University"))
    assert _left_inches(degree) == 0.0
    assert _left_inches(school) == 0.0
