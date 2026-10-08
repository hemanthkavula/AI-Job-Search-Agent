import json
from types import SimpleNamespace

from docx import Document

import app.reference_resume_formatter as formatter
from app.llm_resume_writer import build_prompt, _normalize_skills, _response_request_body
from app.master_resume import fixed_personal_facts, master_resume_payload


EXPECTED_WORD_SHA256 = "2a9d16594d1a618b757b5367c681b389dcf419c8928f95666a0143d4ce019da7"


def test_uploaded_word_is_authoritative_format_record():
    record = formatter.load_word_format()
    assert record["source"]["authority"] == "user_uploaded_word_format_template"
    assert record["source"]["technical_content_authority"] is False
    assert record["source"]["sha256"] == EXPECTED_WORD_SHA256
    assert record["source"]["pages"] == 2
    assert record["style"]["font"] == "Calibri"
    assert record["style"]["name_pt"] == 18
    assert record["style"]["headline_pt"] == 13
    assert record["style"]["contact_pt"] == 11
    assert record["style"]["section_heading_pt"] == 12
    assert record["style"]["summary_pt"] == 11
    assert record["style"]["body_pt"] == 10.0
    assert record["style"]["company_pt"] == 12
    assert record["style"]["job_title_pt"] == 11
    assert record["style"]["education_degree_pt"] == 11
    assert record["style"]["top_margin_in"] == 0.50
    assert record["style"]["bottom_margin_in"] == 0.50
    assert record["style"]["left_margin_in"] == 0.50
    assert record["style"]["right_margin_in"] == 0.50
    assert record["style"]["natural_pagination"] is True
    assert record["style"]["forced_page_breaks"] is False


def test_fixed_personal_facts_exclude_master_technical_content():
    facts = fixed_personal_facts()
    serialized = json.dumps(facts).lower()
    for forbidden in (
        "pyspark",
        "snowflake",
        "databricks",
        "kafka",
        "terraform",
        "great expectations",
        "environment",
        "bullets",
    ):
        assert forbidden not in serialized
    assert facts["name"] == "Hemanth Kavula"
    assert [x["company"] for x in facts["employment_history"]] == [
        "Fidelity Investments",
        "Cigna Healthcare",
        "Target Corporation",
    ]


def test_gpt_json_object_request_explicitly_mentions_json(monkeypatch):
    monkeypatch.setenv("RESUME_LLM_REASONING_EFFORT", "medium")
    monkeypatch.setenv("RESUME_LLM_SERVICE_TIER", "flex")
    monkeypatch.setenv("RESUME_LLM_MAX_OUTPUT_TOKENS", "8000")
    payload = json.loads(
        _response_request_body(
            "gpt-6.1-sol",
            {"task": "Create a tailored resume", "job": {"title": "Data Engineer"}},
        ).decode("utf-8")
    )
    assert payload["model"] == "gpt-6.1-sol"
    assert payload["text"]["format"]["type"] == "json_object"
    assert "json" in payload["input"].lower()
    assert payload["reasoning"]["effort"] == "medium"
    assert payload["service_tier"] == "flex"
    assert payload["max_output_tokens"] == 8000


def test_tailoring_prompt_uses_master_historical_baseline_and_fidelity_jd():
    job = SimpleNamespace(
        company="Example",
        title="Data Engineer",
        description=(
            "Build Python and SQL pipelines on GCP with BigQuery, Dataflow, Pub/Sub and Airflow. "
            "Required hands-on Python, SQL, GCP, BigQuery and Dataflow."
        ),
        description_complete=True,
        description_usable=True,
        tailoring_mode="FULL_JD",
    )
    profile_only = {
        "skills": ["PROFILE_ONLY_TOOL"],
        "summary_source": ["PROFILE ONLY SUMMARY"],
        "experience": [{"evidence": ["PROFILE ONLY BULLET"]}],
    }
    prompt = build_prompt(
        job,
        profile_only,
        coverage_plan={
            "target_count": 5,
            "must_cover_terms": ["Python", "SQL", "GCP", "BigQuery", "Dataflow"],
            "targeted_terms": ["Python", "SQL", "GCP", "BigQuery", "Dataflow"],
            "requirements": [
                {"term": "Python", "classification": "required"},
                {"term": "SQL", "classification": "required"},
                {"term": "GCP", "classification": "required"},
                {"term": "BigQuery", "classification": "required"},
                {"term": "Dataflow", "classification": "required"},
            ],
        },
    )
    serialized = json.dumps(prompt)
    assert "PROFILE_ONLY_TOOL" not in serialized
    assert "PROFILE ONLY SUMMARY" not in serialized
    assert "PROFILE ONLY BULLET" not in serialized

    baseline = prompt["master_resume_technical_baseline"]
    assert baseline["skills"]["Cloud Platforms (AWS)"]
    assert baseline["skills"]["Cloud Platforms (Azure)"]
    assert prompt["employer_cloud_credibility_policy"]["selected_cloud_by_employer"]["Fidelity Investments"] == "GCP"
    assert prompt["employer_cloud_credibility_policy"]["selected_cloud_by_employer"]["Cigna Healthcare"] == "AZURE"
    assert prompt["employer_cloud_credibility_policy"]["selected_cloud_by_employer"]["Target Corporation"] == "AWS"
    assert prompt["historical_timeline_policy"]["forbid_ai_era_technology_in_cigna_and_target"] is True
    assert prompt["historical_timeline_policy"]["older_employer_new_technology_requires_master_backing"] is True
    assert prompt["historical_timeline_policy"]["reject_technology_that_is_not_plausible_for_employer_period"] is True
    assert prompt["skills_policy"]["category_names_are_jd_adaptive"] is True
    assert prompt["skills_policy"]["allow_category_rename_merge_split_reorder"] is True
    assert prompt["skills_policy"]["allow_new_categories_when_jd_supported"] is True
    assert prompt["skills_policy"]["technical_skills_must_reflect_entire_final_resume"] is True
    assert prompt["skills_policy"]["preserve_historical_cloud_skills_from_all_employers"] is True
    assert prompt["skills_policy"]["build_skills_after_final_experience_reconciliation"] is True
    assert prompt["skills_policy"]["add_gcp_category_when_fidelity_selects_gcp"] is True
    assert prompt["structure_contract"]["summary_paragraphs"] == 1
    assert prompt["structure_contract"]["summary_min_words"] == 95
    assert prompt["structure_contract"]["summary_target_words"] == "100-140"
    assert prompt["structure_contract"]["summary_min_sentences"] == 4
    assert prompt["structure_contract"]["employer_footer_label"] == "Environment"
    assert prompt["structure_contract"]["environment_paragraphs"] is False


def test_skill_category_names_can_be_jd_adaptive_and_preserve_all_historical_clouds():
    source = {
        "Programming & Query Languages": ["Python", "SQL", "Java"],
        "Data Processing & Lakehouse": ["Apache Spark", "Databricks", "Delta Lake"],
        "Orchestration & Transformation": ["Airflow", "dbt"],
        "AWS Data Platform": ["AWS Glue", "Amazon S3"],
        "Azure Data Platform": ["Azure Data Factory", "ADLS Gen2"],
    }
    normalized = _normalize_skills(
        source,
        "Build Python, SQL, Java, Spark, Databricks, Airflow and dbt pipelines on AWS using Glue and S3.",
    )
    assert "Programming & Query Languages" in normalized
    assert "Data Processing & Lakehouse" in normalized
    assert "Orchestration & Transformation" in normalized
    assert "AWS Data Platform" in normalized
    assert "Azure Data Platform" in normalized
    assert "Programming Languages" not in normalized
    assert "Cloud Platforms (AWS)" not in normalized
    assert "Cloud Platforms (Azure)" not in normalized
    assert "Java" in normalized["Programming & Query Languages"]
    assert "AWS Glue" in normalized["AWS Data Platform"]
    assert "Azure Data Factory" in normalized["Azure Data Platform"]
    assert "ADLS Gen2" in normalized["Azure Data Platform"]


def test_aws_jd_does_not_remove_cigna_azure_from_technical_skills():
    source = {
        "Programming & Query Languages": ["Python", "SQL"],
        "AWS Data Platform": ["AWS Glue", "Amazon S3"],
    }
    normalized = _normalize_skills(
        source,
        "Required AWS Glue, S3, Python and SQL.",
    )
    all_skills = [value for values in normalized.values() for value in values]
    assert "AWS Glue" in all_skills
    assert "Amazon S3" in all_skills
    assert "Azure Data Factory" in all_skills
    assert "ADLS Gen2" in all_skills


def test_new_jd_supported_skill_category_is_preserved():
    normalized = _normalize_skills(
        {
            "Programming Languages": ["Python", "SQL"],
            "Cloud Platforms (AWS)": ["AWS Glue"],
            "Cloud Platforms (Azure)": ["Azure Data Factory"],
            "Data Observability & Reliability": ["Great Expectations"],
        },
        "Build reliable Python SQL data pipelines with Great Expectations.",
    )
    assert normalized["Data Observability & Reliability"] == ["Great Expectations"]


def test_master_payload_remains_conservative_fallback():
    payload = master_resume_payload()
    assert payload["_master_mode"] is True
    assert len(payload["experience"][0]["bullets"]) == 8
    assert len(payload["experience"][1]["bullets"]) == 7
    assert len(payload["experience"][2]["bullets"]) == 7
    assert "dbt transformation models" in payload["experience"][0]["bullets"][4]
    assert "ML-ready historical and near-real-time datasets" in payload["experience"][0]["bullets"][6]


def _assert_uploaded_master_headers(doc):
    expected = {
        "Fidelity Investments": ("Jersey City, NJ", "Jan 2025 – Present", "Senior Data Engineer"),
        "Cigna Healthcare": ("Bangalore, India", "Jan 2022 – Dec 2023", "Data Engineer"),
        "Target Corporation": ("Bangalore, India", "Jan 2020 – Dec 2021", "Data Engineer"),
    }
    for company, (location, dates, title) in expected.items():
        index = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith(company))
        header = doc.paragraphs[index]
        assert header.text == f"{company} | {location}\t{dates}"
        assert doc.paragraphs[index + 1].text == title
    education = next(p for p in doc.paragraphs if p.text.startswith("Rowan University"))
    assert education.text == "Rowan University | Glassboro, NJ\tJan 2024 – Dec 2025"


def _assert_selective_bold_and_skill_headings_only(doc):
    summary_start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().upper() == "PROFESSIONAL SUMMARY")
    skills_start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().upper() == "TECHNICAL SKILLS")
    summary = [p for p in doc.paragraphs[summary_start + 1:skills_start] if p.text.strip()]
    assert len(summary) == 1
    visible = [r for r in summary[0].runs if r.text]
    assert any(bool(r.bold) for r in visible)
    assert any(not bool(r.bold) for r in visible)

    start = next(i for i, p in enumerate(doc.paragraphs) if p.text == "TECHNICAL SKILLS")
    end = next(i for i, p in enumerate(doc.paragraphs) if p.text == "PROFESSIONAL EXPERIENCE")
    for paragraph in doc.paragraphs[start + 1:end]:
        if not paragraph.text.strip():
            continue
        before, _, after = paragraph.text.partition(":")
        assert before and after
        seen_colon = False
        for run in paragraph.runs:
            text = run.text or ""
            if ":" in text:
                left, right = text.split(":", 1)
                if left:
                    assert bool(run.bold)
                seen_colon = True
                if right.strip():
                    assert not bool(run.bold)
            elif seen_colon and text.strip():
                assert not bool(run.bold)

def test_renderer_reproduces_uploaded_word_format_contract_with_environment_footer(monkeypatch, tmp_path):
    monkeypatch.setattr(formatter, "ROOT", tmp_path)
    monkeypatch.setattr(formatter, "WORD_FORMAT_PATH", formatter.WORD_FORMAT_PATH)
    job = SimpleNamespace(company="Example Company", title="Data Engineer")
    path = formatter.render_llm_resume(
        job, {}, master_resume_payload(), output_dir="resumes"
    )
    result = formatter.validate_master_format_contract(path)
    assert result["passed"], result["reasons"]
    doc = Document(path)
    assert doc.paragraphs[0].text == "Hemanth Kavula"
    assert doc.paragraphs[1].text == "Senior Data Engineer"
    assert sum(p.text == "Roles & Responsibilities:" for p in doc.paragraphs) == 0
    assert sum(p.text.startswith("Environment: ") for p in doc.paragraphs) == 3
    assert sum(p.text.startswith("Skills: ") for p in doc.paragraphs) == 0
    assert not any(p.paragraph_format.page_break_before is True for p in doc.paragraphs)
    _assert_uploaded_master_headers(doc)


def _generated_payload(long_content=False):
    extra = (
        " This additional detail verifies that the renderer keeps the native paragraph formatting even when a bullet wraps."
        if long_content
        else ""
    )
    skills = {
        "Programming Languages": ["Python", "SQL"],
        "Cloud Platforms (AWS)": ["AWS Glue", "Amazon S3"],
        "Cloud Platforms (Azure)": ["Azure Data Factory", "ADLS Gen2"],
        "Data Integration & Orchestration": ["Airflow", "dbt"],
    }
    return {
        "summary": (
            "Senior Data Engineer with more than five years of experience designing scalable batch and streaming data platforms "
            "for enterprise analytics across financial services, healthcare, and retail. Builds reliable Python, SQL, PySpark, "
            "Spark, Kafka, Databricks, Snowflake, AWS, and Azure pipelines with strong attention to production reliability and "
            "maintainable engineering standards. Brings hands-on depth in dimensional modeling, medallion architecture, dbt, "
            "data quality, governance, orchestration, and performance optimization. In the current Fidelity Investments role, "
            "engineers trusted trading and market datasets that support risk, compliance, analytics, and relevant AI/ML use cases. "
            "Partners with technical and business stakeholders to translate complex requirements into governed, scalable data products "
            "while keeping solutions interview-defensible, operationally supportable, and aligned with the target role."
            + extra
        ),
        "skills": skills,
        "experience": [
            {
                "company": "Fidelity Investments",
                "bullets": [
                    f"Built Python and SQL data pipelines for reliable financial data processing workflow {i}.{extra}"
                    for i in range(1, 9)
                ],
                "skills_used": ["Python", "SQL"],
            },
            {
                "company": "Cigna Healthcare",
                "bullets": [
                    f"Developed Python and SQL data integrations for healthcare workflow {i}.{extra}"
                    for i in range(1, 8)
                ],
                "skills_used": ["Python", "SQL"],
            },
            {
                "company": "Target Corporation",
                "bullets": [
                    f"Implemented Python and SQL batch processing for retail workflow {i}.{extra}"
                    for i in range(1, 8)
                ],
                "skills_used": ["Python", "SQL"],
            },
        ],
    }


def test_jd_tailored_renderer_keeps_word_format_and_environment_footers(monkeypatch, tmp_path):
    monkeypatch.setattr(formatter, "ROOT", tmp_path)
    monkeypatch.setattr(formatter, "WORD_FORMAT_PATH", formatter.WORD_FORMAT_PATH)
    job = SimpleNamespace(company="Example Company", title="Senior Data Engineer")
    path = formatter.render_llm_resume(job, {}, _generated_payload(), output_dir="resumes")
    result = formatter.validate_master_format_contract(path)
    assert result["passed"], result["reasons"]
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Environment: Python, SQL" in text
    assert "Skills: Python, SQL" not in text
    assert not any(p.paragraph_format.page_break_before is True for p in doc.paragraphs)
    _assert_uploaded_master_headers(doc)
    _assert_selective_bold_and_skill_headings_only(doc)


def test_renderer_preserves_master_paragraph_format_when_content_wraps(monkeypatch, tmp_path):
    monkeypatch.setattr(formatter, "ROOT", tmp_path)
    monkeypatch.setattr(formatter, "WORD_FORMAT_PATH", formatter.WORD_FORMAT_PATH)
    job = SimpleNamespace(company="Example Company", title="Senior Data Engineer")
    path = formatter.render_llm_resume(job, {}, _generated_payload(long_content=True), output_dir="resumes")
    result = formatter.validate_master_format_contract(path)
    assert result["passed"], result["reasons"]
    assert result["content_length_policy"] == "compact_master_like_layout"
    doc = Document(path)
    _assert_uploaded_master_headers(doc)
