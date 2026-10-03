from types import SimpleNamespace
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from app.reference_resume_formatter import render_llm_resume


def test_reference_layout_uses_calibri_and_keeps_first_bullet_with_employer(tmp_path):
    profile={
      "name":"Candidate","contact":{},
      "experience":[{"company":"Fidelity Investments","location":"Jersey City, NJ","dates":"Jan 2025 – Present","title":"Senior Data Engineer"}],
      "education":[{"degree":"MS Computer Science","school":"Rowan University","location":"Glassboro, NJ","start":"2024","end":"2025"}],
      "output":{"resume_filename_pattern":"Candidate_{Company}_{JobTitle}"}
    }
    generated={"summary":"Data engineer.","skills":{"Programming":["Python","SQL"]},"experience":[{"company":"Fidelity Investments","bullets":["Built reliable data pipelines with Python and SQL.","Optimized Spark workloads."]}]}
    path=render_llm_resume(SimpleNamespace(company="Example",title="Senior Data Engineer"),profile,generated,output_dir=str(tmp_path))
    doc=Document(path)
    assert doc.styles["Normal"].font.name=="Calibri"
    paragraphs=doc.paragraphs
    idx=next(i for i,p in enumerate(paragraphs) if p.text.startswith("Fidelity Investments"))
    company=paragraphs[idx]; title=paragraphs[idx+1]; first=paragraphs[idx+2]
    assert company.paragraph_format.keep_with_next is True
    assert title.paragraph_format.keep_with_next is True
    assert first.paragraph_format.keep_together is True
    assert first.paragraph_format.keep_with_next is not True
    assert "	" in company.text
    assert any(t.alignment==WD_TAB_ALIGNMENT.RIGHT for t in company.paragraph_format.tab_stops)
    section=next(p for p in doc.paragraphs if p.text=="PROFESSIONAL EXPERIENCE")
    assert section.paragraph_format.keep_with_next is True
    assert section.runs[0].font.name=="Calibri"
